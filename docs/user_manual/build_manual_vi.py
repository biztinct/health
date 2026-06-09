#!/usr/bin/env python3
"""Build the Vietnamese Viet UC CMS user manual from content_vi/*.md."""

import glob
import os
import re
import struct

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
CONTENT = os.path.join(HERE, "content_vi")
OUT = os.path.join(HERE, "VietUC_CMS_User_Manual_Vietnamese.docx")

BRAND = RGBColor(0x1A, 0x23, 0x7E)
ACCENT = RGBColor(0xFB, 0x8C, 0x00)
CALLOUT = {
    "Mẹo": (RGBColor(0x1B, 0x5E, 0x20), "E8F5E9"),
    "Lưu ý": (RGBColor(0x0D, 0x47, 0xA1), "E3F2FD"),
    "Cảnh báo": (RGBColor(0xB7, 0x1C, 0x1C), "FFEBEE"),
}


def png_size(path):
    try:
        with open(path, "rb") as file_handle:
            head = file_handle.read(24)
        if head[:8] == b"\x89PNG\r\n\x1a\n":
            return struct.unpack(">II", head[16:24])
    except Exception:
        pass
    return None


def set_cell_shading(cell, fill_hex):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def prevent_row_split(row):
    row_properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    row_properties.append(cant_split)


def create_decimal_numbering(doc):
    numbering_part = doc.part.numbering_part
    numbering = numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    number_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    new_abstract_id = max(abstract_ids) + 1 if abstract_ids else 0
    new_number_id = max(number_ids) + 1 if number_ids else 1

    abstract_number = OxmlElement("w:abstractNum")
    abstract_number.set(qn("w:abstractNumId"), str(new_abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in [
        ("w:start", "1"),
        ("w:numFmt", "decimal"),
        ("w:lvlText", "%1."),
        ("w:lvlJc", "left"),
    ]:
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        level.append(element)
    paragraph_properties = OxmlElement("w:pPr")
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "720")
    indentation.set(qn("w:hanging"), "360")
    paragraph_properties.append(indentation)
    level.append(paragraph_properties)
    abstract_number.append(level)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(new_number_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(new_abstract_id))
    number.append(abstract_reference)
    numbering.insert(0, abstract_number)
    numbering.append(number)
    return new_number_id


def apply_numbering(paragraph, number_id):
    paragraph_properties = paragraph._p.get_or_add_pPr()
    number_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number_properties.append(level)
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(number_id))
    number_properties.append(number)
    paragraph_properties.append(number_properties)


def add_runs(paragraph, text):
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


def add_toc(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fields = [
        ("begin", None),
        ("instr", 'TOC \\o "1-3" \\h \\z \\u'),
        ("separate", None),
        ("text", "Cập nhật trường này (chọn tất cả, nhấn F9) để tạo mục lục."),
        ("end", None),
    ]
    for kind, text in fields:
        if kind == "instr":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        elif kind == "text":
            element = OxmlElement("w:t")
            element.text = text
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        run._r.append(element)


def add_page_number_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Viet UC CMS — Hướng dẫn sử dụng    |    Trang ")
    for kind, text in [("begin", None), ("instr", "PAGE"), ("end", None)]:
        if kind == "instr":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(element)


def add_image(doc, filename, caption):
    path = os.path.join(IMG, filename)
    if not os.path.exists(path):
        doc.add_paragraph("[thiếu ảnh chụp màn hình: %s]" % filename)
        return
    size = png_size(path)
    width = Inches(6.2)
    if size and size[1] and size[0] / size[1] < 0.9:
        width = Inches(2.6)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(path, width=width)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run("Hình: " + caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_callout(doc, kind, text):
    color, fill = CALLOUT[kind]
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    label = paragraph.add_run(kind.upper() + "  ")
    label.bold = True
    label.font.color.rgb = color
    add_runs(paragraph, text)
    doc.add_paragraph()


def add_md_table(doc, rows):
    column_count = max(len(row) for row in rows)
    keep_compact_table_together = len(rows) <= 4
    table = doc.add_table(rows=0, cols=column_count)
    table.style = "Light Grid Accent 1"
    for row_index, row in enumerate(rows):
        table_row = table.add_row()
        prevent_row_split(table_row)
        cells = table_row.cells
        for column_index in range(column_count):
            text = row[column_index] if column_index < len(row) else ""
            cell = cells[column_index]
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], text)
            if keep_compact_table_together and row_index < len(rows) - 1:
                cell.paragraphs[0].paragraph_format.keep_with_next = True
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


IMG_RE = re.compile(r"^!\[(.*?)\]\(IMG:(.+?)\)\s*$")
CALLOUT_RE = re.compile(r"^\*\*(Mẹo|Lưu ý|Cảnh báo):\*\*\s*(.*)$")


def render_markdown(doc, markdown):
    lines = markdown.split("\n")
    index = 0
    current_number_id = None
    while index < len(lines):
        line = lines[index].rstrip()
        is_numbered = bool(re.match(r"^\s*\d+[.)]\s+", line))
        if not is_numbered:
            current_number_id = None
        if not line.strip():
            index += 1
            continue
        match = IMG_RE.match(line)
        if match:
            add_image(doc, match.group(2).strip(), match.group(1).strip() or match.group(2))
            index += 1
            continue
        match = CALLOUT_RE.match(line)
        if match:
            add_callout(doc, match.group(1), match.group(2))
            index += 1
            continue
        if (
            line.lstrip().startswith("|")
            and index + 1 < len(lines)
            and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[index + 1])
        ):
            rows = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                raw = lines[index].strip().strip("|")
                if re.match(r"^[\s:|-]+$", raw):
                    index += 1
                    continue
                rows.append([cell.strip() for cell in raw.split("|")])
                index += 1
            add_md_table(doc, rows)
            continue
        if line.startswith("# "):
            if getattr(doc, "_seen_h1", False):
                doc.add_page_break()
            doc._seen_h1 = True
            doc.add_heading(line[2:].strip(), level=1)
            index += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            index += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            index += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
            index += 1
            continue
        match = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if match:
            if current_number_id is None:
                current_number_id = create_decimal_numbering(doc)
            paragraph = doc.add_paragraph(style="List Paragraph")
            apply_numbering(paragraph, current_number_id)
            add_runs(paragraph, match.group(1))
            index += 1
            continue
        match = re.match(r"^\s*[-*]\s+(.*)$", line)
        if match:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_runs(paragraph, match.group(1))
            index += 1
            continue
        paragraph = doc.add_paragraph()
        add_runs(paragraph, line)
        index += 1


def build():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for level, size, color in [
        (1, 20, BRAND),
        (2, 15, BRAND),
        (3, 12.5, ACCENT),
    ]:
        style = doc.styles["Heading %d" % level]
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("VIET UC CMS")
    run.bold = True
    run.font.size = Pt(40)
    run.font.color.rgb = BRAND
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Hướng dẫn sử dụng")
    run.font.size = Pt(26)
    run.font.color.rgb = ACCENT
    description = doc.add_paragraph()
    description.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = description.add_run("Nền tảng quản lý hoạt động chăm sóc sức khỏe tại nhà")
    run.font.size = Pt(13)
    run.italic = True
    for _ in range(8):
        doc.add_paragraph()
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run(
        "Hướng dẫn đầy đủ cho mọi menu và quy trình — web và di động\n"
    ).italic = True
    metadata.add_run("Phiên bản 1.0").bold = True
    doc.add_page_break()

    toc_title = doc.add_paragraph()
    run = toc_title.add_run("Mục lục")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = BRAND
    doc._seen_h1 = True
    note = doc.add_paragraph()
    add_runs(
        note,
        '**Lưu ý:** Đây là tài liệu Word có thể chỉnh sửa. Để cập nhật số trang bên dưới, chọn toàn bộ tài liệu (Ctrl+A / Cmd+A), nhấn F9, sau đó chọn "Cập nhật toàn bộ bảng". Bản PDF đã có sẵn số trang.',
    )
    add_toc(doc)
    doc.add_page_break()
    doc._seen_h1 = False

    files = sorted(glob.glob(os.path.join(CONTENT, "*.md")))
    if not files:
        raise SystemExit("Không tìm thấy tệp content_vi/*.md.")
    for path in files:
        with open(path, encoding="utf-8") as file_handle:
            render_markdown(doc, file_handle.read())

    add_page_number_footer(doc.sections[0])
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)
    doc.save(OUT)
    print("Wrote", OUT)
    return OUT


if __name__ == "__main__":
    build()
