#!/usr/bin/env python3
"""Local pre-processor (run on the analyst machine, NOT the server).

Reads the legacy .xlsx exports and produces:
  1. <name>.json        — list-of-dicts (original column keys) for the server-side
                          migration.runner (decouples the server from openpyxl).
  2. <name>_EN.xlsx     — same data with English column headers, for human review.

Usage:
  python3 convert_xlsx.py Contact_mig.xlsx Booking_mig.xlsx --out /tmp
"""
import argparse
import json
import os
from datetime import datetime, date, time

import openpyxl

# Vietnamese -> English header translations (Deliverable 1)
HEADER_MAP = {
    # --- Contact_mig ---
    'STT': 'Row No', 'ID': 'Legacy Contact GUID', 'ngay_dau_su_dung': 'First Contact Date',
    'gioi_tinh': 'Gender', 'nam_sinh': 'Date of Birth', 'nguon_khach_hang': 'Lead Source',
    'phong_kham': 'Clinic/Branch', 'customer_id': 'Customer ID', 'ly_do_tu_choi': 'Rejection Reason',
    'so_dien_thoai_2': 'Phone 2', 'NextContactAt': 'Next Contact At', 'so_cccd': 'National ID (CCCD)',
    'ten_benh_nhan': 'Patient Name', 'khoang_cach': 'Distance', 'dich_vu_quan_tam': 'Service of Interest',
    'loai_khach_hnagf': 'Customer Type', 'ghi_chu': 'Note', 'cskh_3': 'Customer Care',
    'phan_hoi_kh': 'Customer Feedback', 'pancake_tag': 'Pancake Tag', 'lien_ket_lich_hen': 'Linked Appointment',
    'FullAddress': 'Full Address', 'country': 'Country', 'province': 'Province', 'district': 'District',
    'commune': 'Commune/Ward', 'kanban_order': 'Kanban Order', 'dich_vu_don_vi_tinh': 'Service Unit',
    'dich_vu_danh_muc': 'Service Category', 'CreatedOn': 'Created On', 'dich_vu_RetailPrice': 'Service Retail Price',
    'dich_vu': 'Services', 'SubTotal': 'Subtotal', 'Discount': 'Discount', 'GrandTotal': 'Grand Total',
    'CreatedBy': 'Created By', 'ModifiedOn': 'Modified On', 'psid': 'PSID', 'TotalQuantity': 'Total Quantity',
    'ModifiedBy': 'Modified By', 'page_id': 'Page ID', 'LastContactUser': 'Last Contact User',
    'ad_id': 'Ad ID', 'TransferMoney': 'Transfer Amount', 'Cash': 'Cash', 'LastContactAt': 'Last Contact At',
    'ad_id_tiktok': 'TikTok Ad ID', 'ad_id_fb': 'Facebook Ad ID',
    # --- Booking_mig ---
    'dich_vu_tai': 'Service Location', 'lien_ket_lien_he': 'Linked Contact',
    'Contact_ad_id_tiktok': 'Contact TikTok Ad ID', 'Contact_dich_vu_danh_muc': 'Contact Service Category',
    'Contact_dich_vu_quan_tam': 'Contact Service of Interest', 'so_dien_thoai': 'Phone',
    'gui_xe': 'Parking Fee', 'Contact_nam_sinh': 'Date of Birth', 'nghe_nghiep': 'Profession',
    'dan_toc': 'Ethnicity', 'Contact_gioi_tinh': 'Gender', 'Client ID': 'Legacy Client Code',
    'SourceofClient': 'Client Source', 'Client_Type': 'Client Type', 'lien_ket_benh_nhan': 'Linked Patient',
    'trang_thai': 'Status', 'ly_do_huy': 'Cancellation Reason', 'Contact_phong_kham': 'Contact Clinic',
    'phong_kham_city': 'Clinic City', 'phu_phi': 'Surcharge', 'nguoi_thuc_hien': 'Performed By',
    'Contact_FullAddress': 'Full Address', 'Contact_Address': 'Address', 'note': 'Note',
    'Contact_customer_id': 'Contact Customer ID', 'phong_kham_Owner': 'Clinic Owner', 'loai_kham': 'Visit Type',
    'time_in': 'Start Time', 'time_out': 'End Time', 'thanh_toan_truoc': 'Prepayment',
    'thanh_toan': 'Payment Status', 'thoi_gian_bat_dau': 'Start Datetime', 'bac_si_dieu_duong': 'Nurse/Doctor',
    'tao_don': 'Order Created', 'so_don_da_tao': 'Orders Created Count', 'hinh_thuc_thanh_toan': 'Payment Method',
    'tra': 'Paid', 'nguoi_phu_trach': 'Person in Charge', 'in_don': 'Printed Order',
    'mo_ta_tinh_trang': 'Condition Description', 'conversation_id': 'Conversation ID',
}


def cell_to_json(v):
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime('%H:%M')
    return v


def convert(path, outdir):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    raw_hdr = list(rows[0])
    # disambiguate duplicate headers
    seen, hdr = {}, []
    for h in raw_hdr:
        h = h if h is not None else 'col'
        if h in seen:
            seen[h] += 1
            hdr.append('%s_%d' % (h, seen[h]))
        else:
            seen[h] = 0
            hdr.append(h)

    records = []
    for r in rows[1:]:
        if all(c in (None, '') for c in r):
            continue
        records.append({hdr[i]: cell_to_json(r[i]) for i in range(len(hdr))})

    base = os.path.splitext(os.path.basename(path))[0]
    json_path = os.path.join(outdir, base + '.json')
    with open(json_path, 'w') as f:
        json.dump(records, f, ensure_ascii=False)

    # English-headed review workbook
    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.append([HEADER_MAP.get(h, h) for h in hdr])
    for rec in records:
        out_ws.append([rec.get(h) for h in hdr])
    en_path = os.path.join(outdir, base + '_EN.xlsx')
    out_wb.save(en_path)
    print('%s -> %s (%d rows), %s' % (os.path.basename(path), json_path, len(records), en_path))
    return json_path


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('files', nargs='+')
    ap.add_argument('--out', default='/tmp')
    a = ap.parse_args()
    for f in a.files:
        convert(f, a.out)
