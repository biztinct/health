# -*- coding: utf-8 -*-
"""ACCESS REVAMP AR-1 — who may give what, several at once, on somebody's behalf.

The numbers in the class names are the handover's test numbers
(`docs/handovers/ACCESS_AR1_REPAIR_POLICY_PEOPLE.md` §3.K), so a report can say
"3 passed" and mean one thing.

EVERY CALL HERE IS MADE AS A REAL PERSON. `self.env` in a test is the
superuser, and the superuser is exempt from both new rules on purpose (it is
how a customer's first administrator is created). A test of "only holders may
give it" run as the superuser would pass for the wrong reason, so every refusal
below goes through `with_user(...)`.
"""

import base64
import io
from datetime import timedelta
from unittest.mock import patch

from odoo import fields
from odoo.exceptions import AccessError, UserError
from odoo.tests import TransactionCase, tagged

from odoo.addons.biz_access.models.access_common import (
    HIDDEN_LOGINS_PARAM, SYSTEM_APPLY_KEY, forbidden_in_closure, hidden_logins,
    user_can_manage)


class Ar1Case(TransactionCase):
    """Three throwaway permissions, three roles over them, and the people who
    give and hold them — built fresh so nothing depends on how a database is
    gated today."""

    def setUp(self):
        super().setUp()
        stamp = str(fields.Datetime.now()).replace(' ', '').replace(
            ':', '').replace('-', '')
        self.stamp = stamp
        self.Users = self.env['res.users'].with_context(no_reset_password=True)
        self.team = self.env.ref('biz_access.group_access_manager')

        self.g_a = self.env['res.groups'].create({'name': 'ZZ AR1 A %s' % stamp})
        self.g_b = self.env['res.groups'].create({'name': 'ZZ AR1 B %s' % stamp})
        self.g_c = self.env['res.groups'].create({'name': 'ZZ AR1 C %s' % stamp})
        self.g_own = self.env['res.groups'].create(
            {'name': 'ZZ AR1 Owner %s' % stamp})

        self.role_a = self._role('ZZ AR1 Role A %s' % stamp, self.g_a)
        self.role_b = self._role('ZZ AR1 Role B %s' % stamp, self.g_b)
        self.role_c = self._role('ZZ AR1 Role C %s' % stamp, self.g_c)
        self.owner_role = self._role('ZZ AR1 Owner %s' % stamp, self.g_own,
                                     guarded=True)

        self.manager = self._user('mgr', self.team)
        self.owner = self._user('own', self.team | self.g_own)
        self.lan = self._user('lan')
        self.mai = self._user('mai')
        self.plain = self._user('plain')
        self.sysadmin = self._user('sys', self.env.ref('base.group_system'))

    # ----------------------------------------------------------- fixtures
    def _role(self, name, groups, guarded=False):
        ability = self.env['biz.access.ability'].create({
            'technical_key': 'zz-ar1-%s' % name.lower().replace(' ', '-'),
            'name': name, 'area': 'general',
            'description': 'A throwaway ability for a test.',
            'group_ids': [(6, 0, groups.ids)],
        })
        return self.env['biz.access.role'].create({
            'name': name, 'area': 'general', 'guarded': guarded,
            'description': 'A throwaway role for a test.',
            'ability_ids': [(6, 0, ability.ids)],
        })

    def _user(self, tag, groups=None, name=None):
        user = self.Users.create({
            'name': name or 'ZZ AR1 %s %s' % (tag.capitalize(), self.stamp),
            'login': 'zz.ar1.%s.%s@example.com' % (tag, self.stamp),
            'email': 'zz.ar1.%s.%s@example.com' % (tag, self.stamp),
            'group_ids': [(4, self.env.ref('base.group_user').id)],
        })
        if groups:
            user.write({'group_ids': [(4, g.id) for g in groups]})
        return user

    def as_(self, user):
        return self.env['biz.access'].with_user(user)

    def holds(self, user, role):
        user.invalidate_recordset(['all_group_ids', 'group_ids'])
        return set(role.group_ids.ids) <= set(user.all_group_ids.ids)

    def audit_rows(self, user, origin='board'):
        return self.env['biz.access.delegation'].search(
            [('delegate_user_id', '=', user.id), ('origin', '=', origin)])


# =========================================================================
#  3 — a guarded role is given (or lent) only by its holders
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_03_Guarded(Ar1Case):

    def test_a_manager_who_does_not_hold_it_cannot_give_it(self):
        with self.assertRaises(UserError) as caught:
            self.as_(self.manager).grant(self.owner_role.id, self.lan.id)
        self.assertIn(self.owner_role.name, str(caught.exception))
        self.assertIn('Only somebody who holds', str(caught.exception))
        self.assertFalse(self.holds(self.lan, self.owner_role))

    def test_nor_among_several(self):
        with self.assertRaises(UserError) as caught:
            self.as_(self.manager).grant_many(
                [self.role_a.id, self.owner_role.id], self.lan.id)
        self.assertIn(self.owner_role.name, str(caught.exception))
        # All or nothing: the ordinary role did not slip through either.
        self.assertFalse(self.holds(self.lan, self.role_a))

    def test_nor_by_copying_it_from_somebody_who_has_it(self):
        self.env['biz.access'].grant_many(
            [self.role_a.id, self.owner_role.id], self.mai.id)
        res = self.as_(self.manager).copy_roles(self.mai.id, self.lan.id)
        self.assertTrue(self.holds(self.lan, self.role_a))
        self.assertFalse(self.holds(self.lan, self.owner_role))
        self.assertIn('left out', res['message'])
        self.assertIn(self.owner_role.name, res['message'])

    def test_nor_by_lending_it_on_the_holders_behalf(self):
        with self.assertRaises(UserError) as caught:
            self.as_(self.manager).delegate({
                'delegator_user_id': self.owner.id,
                'delegate_user_id': self.lan.id,
                'profile_ids': [self.owner_role.id],
                'kind': 'permanent'})
        self.assertIn(self.owner_role.name, str(caught.exception))
        self.assertIn('lend', str(caught.exception))

    def test_somebody_who_holds_it_can(self):
        res = self.as_(self.owner).grant(self.owner_role.id, self.lan.id)
        self.assertTrue(res['ok'])
        self.assertTrue(self.holds(self.lan, self.owner_role))

    def test_the_system_administrator_is_exempt(self):
        res = self.as_(self.sysadmin).grant(self.owner_role.id, self.lan.id)
        self.assertTrue(res['ok'])

    def test_the_guard_cannot_be_unticked_by_a_non_holder(self):
        with self.assertRaises(UserError):
            self.owner_role.with_user(self.manager).write({'guarded': False})
        # Nor widened from the ability's side.
        ability = self.owner_role.ability_ids[:1]
        with self.assertRaises(UserError):
            ability.with_user(self.manager).write(
                {'group_ids': [(4, self.g_a.id)]})
        # The order and the sentence stay anybody's to change.
        self.owner_role.with_user(self.manager).write({'sequence': 77})
        self.owner_role.with_user(self.owner).write({'guarded': False})
        self.assertFalse(self.owner_role.guarded)

    def test_the_board_says_which_roles_the_reader_may_give(self):
        board = self.as_(self.manager).get_board(search=self.stamp)
        rows = {p['id']: p for p in board['profiles']}
        self.assertTrue(rows[self.owner_role.id]['guarded'])
        self.assertFalse(rows[self.owner_role.id]['can_give'])
        self.assertTrue(rows[self.role_a.id]['can_give'])
        mine = {p['id']: p for p in self.as_(self.owner).get_board(
            search=self.stamp)['profiles']}
        self.assertTrue(mine[self.owner_role.id]['can_give'])


# =========================================================================
#  4 — nobody gives a role to themselves
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_04_NotToYourself(Ar1Case):

    def assertRefusedToSelf(self, fn):
        with self.assertRaises(UserError) as caught:
            fn()
        self.assertIn('yourself', str(caught.exception))

    def test_grant_grant_many_and_copy_all_refuse(self):
        mgr = self.as_(self.manager)
        self.assertRefusedToSelf(
            lambda: mgr.grant(self.role_a.id, self.manager.id))
        self.assertRefusedToSelf(
            lambda: mgr.grant_many([self.role_a.id], self.manager.id))
        self.env['biz.access'].grant(self.role_b.id, self.mai.id)
        self.assertRefusedToSelf(
            lambda: mgr.copy_roles(self.mai.id, self.manager.id))
        self.assertFalse(self.holds(self.manager, self.role_a))
        self.assertFalse(self.holds(self.manager, self.role_b))

    def test_sudo_does_not_change_who_is_asking(self):
        """Screens reach the facade through `.sudo()`; the uid is still the
        person, and the rule still binds them."""
        self.assertRefusedToSelf(
            lambda: self.as_(self.manager).sudo().grant(
                self.role_a.id, self.manager.id))

    def test_the_system_apply_key_is_refused_to_an_ordinary_caller(self):
        mgr = self.as_(self.manager).with_context(**{SYSTEM_APPLY_KEY: True})
        with self.assertRaises(AccessError):
            mgr.grant(self.role_a.id, self.manager.id)

    def test_and_honoured_for_the_system_itself(self):
        res = self.as_(self.manager).sudo().with_context(
            **{SYSTEM_APPLY_KEY: True}).grant(self.role_a.id, self.manager.id)
        self.assertTrue(res['ok'])

    def test_provisioning_as_the_superuser_still_grants_a_guarded_role(self):
        """What `biz_tenants` does for a customer's first administrator."""
        res = self.env['biz.access'].sudo().grant(
            self.owner_role.id, self.lan.id, reason='Given when created.')
        self.assertTrue(res['ok'])
        self.assertTrue(self.holds(self.lan, self.owner_role))

    def test_the_picker_still_leaves_you_out(self):
        ids = [r['id'] for r in self.as_(self.manager).user_options('ZZ AR1')]
        self.assertNotIn(self.manager.id, ids)
        ids = [r['id'] for r in
               self.as_(self.manager).user_options('ZZ AR1', True)]
        self.assertIn(self.manager.id, ids)

    def test_the_board_carries_the_sentence_for_your_own_passport(self):
        self.assertIn('yourself',
                      self.as_(self.manager).get_board()['self_grant_note'])


# =========================================================================
#  5 — several roles at once
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_05_GrantMany(Ar1Case):

    def test_three_roles_one_write_one_line_of_history(self):
        Users = type(self.env['res.users'])
        real = Users.write
        writes = []

        def counting(recs, vals):
            if 'group_ids' in vals and self.lan.id in recs.ids:
                writes.append(vals)
            return real(recs, vals)

        with patch.object(Users, 'write', counting):
            res = self.as_(self.manager).grant_many(
                [self.role_a.id, self.role_b.id, self.role_c.id], self.lan.id,
                'Starting on Monday')
        self.assertEqual(len(writes), 1, 'more than one write to the person')
        rows = self.audit_rows(self.lan)
        self.assertEqual(len(rows), 1)
        self.assertEqual(set(rows.profile_ids.ids),
                         {self.role_a.id, self.role_b.id, self.role_c.id})
        self.assertEqual(rows.reason, 'Starting on Monday')
        self.assertEqual(res['count'], 3)
        self.assertIn('3 roles', res['message'])
        for role in (self.role_a, self.role_b, self.role_c):
            self.assertTrue(self.holds(self.lan, role))

    def test_already_held_ones_are_skipped_and_said(self):
        self.env['biz.access'].grant(self.role_a.id, self.lan.id)
        res = self.as_(self.manager).grant_many(
            [self.role_a.id, self.role_b.id], self.lan.id)
        self.assertEqual(res['profile_ids'], [self.role_b.id])
        self.assertIn('already theirs', res['message'])

    def test_all_already_held_is_refused(self):
        self.env['biz.access'].grant_many(
            [self.role_a.id, self.role_b.id], self.lan.id)
        with self.assertRaises(UserError) as caught:
            self.as_(self.manager).grant_many(
                [self.role_a.id, self.role_b.id], self.lan.id)
        self.assertIn('already has all', str(caught.exception))

    def test_capped_at_a_hundred(self):
        with self.assertRaises(UserError) as caught:
            self.as_(self.manager).grant_many(list(range(1, 102)), self.lan.id)
        self.assertIn('100', str(caught.exception))

    def test_somebody_who_cannot_manage_is_refused(self):
        with self.assertRaises(AccessError):
            self.as_(self.plain).grant_many([self.role_a.id], self.lan.id)


# =========================================================================
#  6 — copying from a colleague
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_06_CopyRoles(Ar1Case):

    def test_lent_roles_are_never_copied(self):
        facade = self.env['biz.access']
        facade.grant(self.role_a.id, self.mai.id)
        facade.grant(self.role_b.id, self.owner.id)
        facade.with_user(self.owner).delegate({
            'delegate_user_id': self.mai.id,
            'profile_ids': [self.role_b.id], 'kind': 'temporary',
            'date_end': fields.Date.today() + timedelta(days=5)})
        self.assertTrue(self.holds(self.mai, self.role_b))

        preview = self.as_(self.manager).copy_preview(self.mai.id, self.lan.id)
        self.assertEqual([r['id'] for r in preview['roles']], [self.role_a.id])
        self.assertEqual(preview['lent_skipped'], 1)

        self.as_(self.manager).copy_roles(self.mai.id, self.lan.id)
        self.assertTrue(self.holds(self.lan, self.role_a))
        self.assertFalse(self.holds(self.lan, self.role_b))
        self.assertEqual(len(self.audit_rows(self.lan)), 1)

    def test_a_guarded_role_is_dropped_with_a_note(self):
        self.env['biz.access'].grant_many(
            [self.role_a.id, self.owner_role.id], self.mai.id)
        preview = self.as_(self.manager).copy_preview(self.mai.id, self.lan.id)
        self.assertEqual(preview['dropped'], [self.owner_role.name])
        self.assertIn('left out', preview['dropped_note'])

    def test_nothing_to_copy_is_refused_in_words(self):
        with self.assertRaises(UserError) as caught:
            self.as_(self.manager).copy_roles(self.mai.id, self.lan.id)
        self.assertIn('nothing to copy', str(caught.exception))
        # Only a guarded role to copy is a refusal that says why.
        self.env['biz.access'].grant(self.owner_role.id, self.mai.id)
        with self.assertRaises(UserError) as caught:
            self.as_(self.manager).copy_roles(self.mai.id, self.lan.id)
        self.assertIn('left out', str(caught.exception))

    def test_copying_from_yourself_to_a_colleague_is_ordinary(self):
        self.env['biz.access'].grant(self.role_c.id, self.owner.id)
        res = self.as_(self.owner).copy_roles(self.owner.id, self.lan.id)
        self.assertTrue(res['ok'])
        self.assertTrue(self.holds(self.lan, self.role_c))


# =========================================================================
#  7 — a hand-over started on somebody's behalf
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_07_OnBehalf(Ar1Case):

    def setUp(self):
        super().setUp()
        self.env['biz.access'].grant(self.role_a.id, self.mai.id)   # the nurse

    def test_a_manager_starts_one_for_somebody_who_is_away(self):
        offered = [r['id'] for r in
                   self.as_(self.manager).held_roles_of(self.mai.id)]
        self.assertIn(self.role_a.id, offered)
        self.assertNotIn(self.role_b.id, offered)

        res = self.as_(self.manager).delegate({
            'delegator_user_id': self.mai.id,
            'delegate_user_id': self.lan.id,
            'profile_ids': [self.role_a.id], 'kind': 'temporary',
            'date_end': fields.Date.today() + timedelta(days=7)})
        self.assertTrue(res['ok'])
        self.assertIn(self.mai.name, res['message'])
        rec = self.env['biz.access.delegation'].browse(res['id'])
        self.assertEqual(rec.state, 'active')
        self.assertEqual(rec.delegator_user_id, self.mai)
        self.assertEqual(rec.create_uid, self.manager)
        self.assertIn(self.manager.name, rec.started_by_note)
        row = self.as_(self.manager)._delegation_row(rec)
        self.assertIn('Started by', row['started_by'])
        self.assertTrue(self.holds(self.lan, self.role_a))

    def test_only_what_the_absent_person_holds_can_be_lent(self):
        with self.assertRaises(UserError):
            self.as_(self.manager).delegate({
                'delegator_user_id': self.mai.id,
                'delegate_user_id': self.lan.id,
                'profile_ids': [self.role_b.id], 'kind': 'permanent'})

    def test_somebody_who_cannot_manage_cannot_name_a_lender(self):
        with self.assertRaises(AccessError):
            self.as_(self.plain).delegate({
                'delegator_user_id': self.mai.id,
                'delegate_user_id': self.lan.id,
                'profile_ids': [self.role_a.id], 'kind': 'permanent'})
        with self.assertRaises(AccessError):
            self.as_(self.plain).held_roles_of(self.mai.id)

    def test_your_own_hand_over_carries_no_started_by(self):
        res = self.as_(self.mai).delegate({
            'delegate_user_id': self.lan.id,
            'profile_ids': [self.role_a.id], 'kind': 'permanent'})
        rec = self.env['biz.access.delegation'].browse(res['id'])
        self.assertFalse(rec.started_by_note)


# =========================================================================
#  8 — the default window comes from the setting
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_08_DefaultWindow(Ar1Case):

    def test_the_setting_reaches_the_board(self):
        self.env['ir.config_parameter'].sudo().set_param(
            'biz_access.default_window_days', '7')
        self.assertEqual(
            self.as_(self.plain).get_board()['default_window_days'], 7)

    def test_nonsense_falls_back_and_zero_is_never_offered(self):
        Param = self.env['ir.config_parameter'].sudo()
        Param.set_param('biz_access.default_window_days', 'soon')
        self.assertEqual(self.env['biz.access.delegation'].default_end_days(),
                         14)
        Param.set_param('biz_access.default_window_days', '0')
        self.assertEqual(self.env['biz.access.delegation'].default_end_days(),
                         1)


# =========================================================================
#  9 — the Hand-overs lens: chips and a name search
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_09_HandoverLens(Ar1Case):

    def setUp(self):
        super().setUp()
        facade = self.env['biz.access']
        facade.grant(self.role_a.id, self.mai.id)
        today = fields.Date.today()
        self.running = facade.with_user(self.mai).delegate({
            'delegate_user_id': self.lan.id, 'profile_ids': [self.role_a.id],
            'kind': 'temporary', 'date_end': today + timedelta(days=3)})['id']
        self.back = facade.with_user(self.mai).delegate({
            'delegate_user_id': self.plain.id, 'profile_ids': [self.role_a.id],
            'kind': 'temporary', 'date_end': today + timedelta(days=3)})['id']
        facade.with_user(self.mai).revoke(self.back)

    def ids(self, res):
        return {r['id'] for r in res['rows']}

    def test_the_default_is_running(self):
        board = self.as_(self.manager).get_board()
        self.assertEqual(board['handovers']['state'], 'running')
        ids = {r['id'] for r in board['delegations']}
        self.assertIn(self.running, ids)
        self.assertNotIn(self.back, ids)

    def test_the_chips_filter(self):
        mgr = self.as_(self.manager)
        self.assertNotIn(self.running, self.ids(mgr.handovers('taken_back')))
        self.assertIn(self.back, self.ids(mgr.handovers('taken_back')))
        self.assertTrue({self.running, self.back}
                        <= self.ids(mgr.handovers('all')))
        self.assertNotIn(self.back, self.ids(mgr.handovers('ended')))

    def test_a_name_matches_either_person(self):
        mgr = self.as_(self.manager)
        by_lender = mgr.handovers('all', 'ar1 mai %s' % self.stamp)
        self.assertTrue({self.running, self.back} <= self.ids(by_lender))
        by_borrower = mgr.handovers('all', 'AR1 LAN %s' % self.stamp)
        self.assertEqual(self.ids(by_borrower), {self.running})
        self.assertEqual(by_borrower['counts']['running'], 1)
        self.assertEqual(by_borrower['counts']['taken_back'], 0)
        self.assertIn('1 running', by_borrower['headline'])

    def test_an_ordinary_person_still_sees_only_their_own(self):
        ids = self.ids(self.as_(self.plain).handovers('all'))
        self.assertEqual(ids & {self.running, self.back}, {self.back})


# =========================================================================
#  10 — "people hold one" counts everybody, not the faces
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_10_Counter(Ar1Case):

    def test_forty_five_holders_are_forty_five(self):
        group = self.env['res.groups'].create(
            {'name': 'ZZ AR1 Busy %s' % self.stamp})
        role = self._role('ZZ AR1 Busy role %s' % self.stamp, group)
        users = self.Users.create([{
            'name': 'ZZ AR1 Holder %02d %s' % (i, self.stamp),
            'login': 'zz.ar1.h%02d.%s@example.com' % (i, self.stamp),
            'group_ids': [(4, self.env.ref('base.group_user').id),
                          (4, group.id)],
        } for i in range(45)])
        self.assertEqual(len(users), 45)
        board = self.as_(self.manager).get_board(
            search='ZZ AR1 Busy role %s' % self.stamp)
        self.assertEqual([p['id'] for p in board['profiles']], [role.id])
        self.assertEqual(board['profiles'][0]['holder_count'], 45)
        self.assertEqual(len(board['profiles'][0]['holders']), 40)
        self.assertEqual(board['kpis']['people'], 45)


# =========================================================================
#  11 — an account that is not a person is on no list
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_11_HiddenAccount(Ar1Case):

    def setUp(self):
        super().setUp()
        self.ghost = self._user('ghost', name='ZZ AR1 Ghost %s' % self.stamp)
        self.env['ir.config_parameter'].sudo().set_param(
            HIDDEN_LOGINS_PARAM, self.ghost.login)
        facade = self.env['biz.access']
        facade.grant(self.role_a.id, self.ghost.id)
        facade.grant(self.role_a.id, self.mai.id)
        # Made as the system: a person cannot lend anything TO this account
        # (that is one of the refusals below), but history made before it was
        # hidden, or by the platform, must still stay off every list.
        facade.delegate({
            'delegator_user_id': self.mai.id,
            'delegate_user_id': self.ghost.id, 'profile_ids': [self.role_a.id],
            'kind': 'permanent'})

    def test_and_it_cannot_be_lent_to_either(self):
        with self.assertRaises(UserError):
            self.as_(self.mai).delegate({
                'delegate_user_id': self.ghost.id,
                'profile_ids': [self.role_a.id], 'kind': 'permanent'})

    def test_the_setting_is_read(self):
        self.assertIn(self.ghost.login, hidden_logins(self.env))

    def test_not_in_people_the_picker_or_the_spectacles(self):
        mgr = self.as_(self.manager)
        self.assertNotIn(self.ghost.id, [r['id'] for r in mgr.people('ZZ AR1')])
        self.assertNotIn(self.ghost.id,
                         [r['id'] for r in mgr.user_options('Ghost', True)])
        with self.assertRaises(UserError):
            mgr.as_user(self.ghost.id)
        with self.assertRaises(UserError):
            mgr.passport(self.ghost.id)

    def test_not_among_a_roles_holders_nor_the_hand_overs(self):
        mgr = self.as_(self.manager)
        board = mgr.get_board(search=self.role_a.name)
        row = [p for p in board['profiles'] if p['id'] == self.role_a.id][0]
        self.assertNotIn(self.ghost.id, [h['id'] for h in row['holders']])
        self.assertEqual(row['holder_count'], 1)
        self.assertFalse([r for r in mgr.handovers('all')['rows']
                          if r['to_id'] == self.ghost.id])

    def test_not_in_either_spreadsheet(self):
        import openpyxl
        mgr = self.as_(self.manager)
        for res in (mgr.export_roles(), mgr.export_delegations()):
            wb = openpyxl.load_workbook(io.BytesIO(
                base64.b64decode(res['file_b64'])))
            text = ' '.join(str(c.value) for row in wb.active.iter_rows()
                            for c in row if c.value)
            self.assertNotIn('Ghost', text)

    def test_and_nothing_can_be_given_to_it(self):
        with self.assertRaises(UserError):
            self.as_(self.manager).grant(self.role_b.id, self.ghost.id)

    def test_it_still_sees_itself(self):
        rows = self.env['biz.access'].with_user(self.ghost).people()
        self.assertEqual([r['id'] for r in rows], [self.ghost.id])


# =========================================================================
#  12 — the seams nobody had tested
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_12_UntestedSeams(Ar1Case):

    def test_delegate_end_to_end(self):
        self.env['biz.access'].grant(self.role_a.id, self.mai.id)
        res = self.as_(self.mai).delegate({
            'delegate_user_id': self.lan.id, 'profile_ids': [self.role_a.id],
            'kind': 'temporary',
            'date_end': fields.Date.today() + timedelta(days=4),
            'reason': 'Two weeks away'})
        rec = self.env['biz.access.delegation'].browse(res['id'])
        self.assertEqual(rec.state, 'active')
        self.assertEqual(rec.origin, 'delegation')
        self.assertEqual(rec.applied_group_ids, self.g_a)
        self.assertTrue(rec.applied_on)
        self.assertTrue(self.holds(self.lan, self.role_a))

    def test_the_night_ends_what_is_due_and_leaves_what_is_not(self):
        self.env['biz.access'].grant(self.role_a.id, self.mai.id)
        today = fields.Date.today()
        facade = self.env['biz.access'].with_user(self.mai)
        due = facade.delegate({
            'delegate_user_id': self.lan.id, 'profile_ids': [self.role_a.id],
            'kind': 'temporary', 'date_start': today - timedelta(days=5),
            'date_end': today - timedelta(days=1)})['id']
        running = facade.delegate({
            'delegate_user_id': self.plain.id, 'profile_ids': [self.role_a.id],
            'kind': 'temporary', 'date_end': today + timedelta(days=5)})['id']
        res = self.env['biz.access.delegation'].run_auto_revert()
        Deleg = self.env['biz.access.delegation']
        self.assertEqual(Deleg.browse(due).state, 'expired')
        self.assertEqual(Deleg.browse(running).state, 'active')
        self.assertGreaterEqual(res['ended'], 1)
        self.assertFalse(self.holds(self.lan, self.role_a))
        self.assertTrue(self.holds(self.plain, self.role_a))

    def _headers(self, res):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(
            res['file_b64'])))
        return [c.value for c in next(wb.active.iter_rows(max_row=1))]

    def test_both_spreadsheets_have_their_headers(self):
        mgr = self.as_(self.manager)
        roles = mgr.export_roles()
        self.assertTrue(roles['filename'].endswith('.xlsx'))
        self.assertEqual(self._headers(roles)[:4],
                         ['Role', 'Area', 'What this lets someone do',
                          'Person'])
        history = mgr.export_delegations()
        heads = self._headers(history)
        self.assertEqual(heads[:4], ['When', 'How it happened', 'Lent by',
                                     'Lent to'])
        self.assertIn('Started by', heads)

    def test_the_picker_leaves_out_you_and_portal_accounts(self):
        portal = self.Users.create({
            'name': 'ZZ AR1 Portal %s' % self.stamp,
            'login': 'zz.ar1.portal.%s@example.com' % self.stamp,
            'group_ids': [(6, 0, [self.env.ref('base.group_portal').id])],
        })
        ids = [r['id'] for r in self.as_(self.manager).user_options('ZZ AR1')]
        self.assertNotIn(self.manager.id, ids)
        self.assertNotIn(portal.id, ids)
        self.assertIn(self.lan.id, ids)


# =========================================================================
#  A / 2 — the manage gate is one answer, and the access team is not the keys
# =========================================================================
@tagged('post_install', '-at_install')
class TestAr1_ManageGate(Ar1Case):

    def test_one_answer_for_the_facade_and_the_backstop(self):
        for user in (self.manager, self.owner, self.plain, self.sysadmin):
            self.assertEqual(
                user_can_manage(user),
                self.env['biz.access'].with_user(user).can_manage(),
                user.name)

    def test_the_access_team_carries_nothing_forbidden(self):
        self.assertFalse(forbidden_in_closure(self.team, self.env))
