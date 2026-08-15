# -*- coding: utf-8 -*-
"""Round-trip spreadsheet for the dropdown vocabularies.

THE PROBLEM THIS SOLVES
-----------------------
The twenty-three vocabularies share one table, so a naive "Import records" on
the values list would need a `category` column on every row and the client
would have to know the technical keys. Nobody should have to.

Instead there are two shapes, and the client picks whichever suits:

  * ONE VOCABULARY AT A TIME — Master Data > Value Categories > Manage Values.
    That screen carries `default_category_id`, so Odoo's own importer needs
    only code/name/name_vi columns. No wizard involved.

  * ALL OF THEM AT ONCE — this wizard. Download Template produces a workbook
    with ONE SHEET PER VOCABULARY, pre-filled with what is there today; the
    client edits it and uploads it back. The sheet NAME is the category, so
    there is no category column and nothing to look up.

Matching is on `code` within the sheet's category, so re-uploading the same
file updates rather than duplicating — the classic spreadsheet round-trip
failure. A row with a new code is created; a row whose code is missing from
the sheet is LEFT ALONE (never deleted — deletion stays a deliberate act on
the screen, where the in-use guard can explain itself).

Precedent: advanced_pricing's price-list importer already dispatches on sheet
name, so this is a shape the team knows.
"""
import base64
import io
import logging
import re

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

COLUMNS = ['code', 'name', 'name_vi', 'sequence', 'active']
# Excel sheet names: 31 chars max, and []:*?/\ are illegal.
_ILLEGAL_SHEET = re.compile(r'[\[\]:*?/\\]')


def _sheet_name(label):
    return _ILLEGAL_SHEET.sub(' ', label or '')[:31].strip()


def _normalise(text):
    return re.sub(r'\s+', ' ', (text or '')).strip().lower()


def _load_openpyxl():
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise UserError(_(
            'The Python library "openpyxl" is required to import or export '
            'spreadsheets and is not installed on this server.')) from exc
    return openpyxl


class HealthLookupImportWizard(models.TransientModel):
    _name = 'health.lookup.import.wizard'
    _description = 'Import Dropdown Values'

    file = fields.Binary(string='Spreadsheet', attachment=False)
    filename = fields.Char()
    template_file = fields.Binary(readonly=True, attachment=False)
    template_filename = fields.Char(readonly=True)
    log = fields.Text(readonly=True)
    state = fields.Selection(
        [('draft', 'Draft'), ('done', 'Done')], default='draft')

    # ------------------------------------------------------------------
    # Template
    # ------------------------------------------------------------------
    def action_download_template(self):
        """One sheet per vocabulary, pre-filled with today's values.

        Pre-filling is the point: the client never types a code by hand, so a
        re-upload updates the right rows instead of creating near-duplicates.
        """
        self.ensure_one()
        openpyxl = _load_openpyxl()
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)

        categories = self.env['health.lookup.category'].with_context(
            active_test=False).search([], order='name')
        used = set()
        for category in categories:
            title = _sheet_name(category.name) or category.code
            # Two vocabularies could truncate to the same 31 chars.
            base, n = title, 2
            while title.lower() in used:
                suffix = ' (%s)' % n
                title = base[:31 - len(suffix)] + suffix
                n += 1
            used.add(title.lower())

            sheet = workbook.create_sheet(title=title)
            sheet.append(COLUMNS)
            values = self.env['health.lookup.value'].with_context(
                active_test=False).search(
                    [('category_id', '=', category.id)], order='sequence, name')
            for value in values:
                english = value.with_context(lang='en_US').name or ''
                vietnamese = value.name_vi or ''
                # An untranslated value reads back as the English fallback.
                # Writing that into the name_vi column would tell the client
                # the row is already translated when it is not — leave it
                # blank so the gaps are visible at a glance.
                if vietnamese == english:
                    vietnamese = ''
                sheet.append([
                    value.code, english, vietnamese, value.sequence,
                    'yes' if value.active else 'no',
                ])
            for column, width in zip('ABCDE', (26, 42, 42, 12, 10)):
                sheet.column_dimensions[column].width = width

        stream = io.BytesIO()
        workbook.save(stream)
        self.write({
            'template_file': base64.b64encode(stream.getvalue()),
            'template_filename': 'dropdown_values_template.xlsx',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=%s&id=%s&field=template_file'
                   '&filename_field=template_filename&download=true' % (
                       self._name, self.id),
            'target': 'self',
        }

    # ------------------------------------------------------------------
    # Import
    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_('Choose a spreadsheet to import first.'))
        openpyxl = _load_openpyxl()
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.file)), data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise UserError(_('That file could not be read as .xlsx: %s', exc)) from exc

        Category = self.env['health.lookup.category'].with_context(active_test=False)
        Value = self.env['health.lookup.value'].with_context(active_test=False)
        categories = Category.search([])
        # A sheet may be named after the vocabulary OR its technical code.
        by_key = {}
        for category in categories:
            by_key[_normalise(category.name)] = category
            by_key[_normalise(category.code)] = category
            by_key[_normalise(_sheet_name(category.name))] = category

        lines, created, updated, skipped = [], 0, 0, 0
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            category = by_key.get(_normalise(sheet_name))
            if not category:
                lines.append(_('SKIPPED sheet "%s" — no vocabulary with that '
                               'name or code.') % sheet_name)
                skipped += 1
                continue

            header = [str(c.value or '').strip().lower()
                      for c in next(sheet.iter_rows(min_row=1, max_row=1))]
            if 'code' not in header or 'name' not in header:
                lines.append(_('SKIPPED sheet "%s" — needs at least "code" '
                               'and "name" columns.') % sheet_name)
                skipped += 1
                continue
            index = {name: position for position, name in enumerate(header)}

            existing = {v.code: v for v in Value.search(
                [('category_id', '=', category.id)])}
            sheet_created = sheet_updated = 0
            for row in sheet.iter_rows(min_row=2, values_only=True):
                def cell(column):
                    position = index.get(column)
                    if position is None or position >= len(row):
                        return None
                    value = row[position]
                    return str(value).strip() if value is not None else None

                code = cell('code')
                name = cell('name')
                if not code and not name:
                    continue
                if not code:
                    lines.append(_('  sheet "%(sheet)s": row for "%(name)s" '
                                   'has no code — skipped.',
                                   sheet=sheet_name, name=name))
                    continue

                vals = {}
                if name:
                    vals['name'] = name
                sequence = cell('sequence')
                if sequence and sequence.replace('.', '', 1).isdigit():
                    vals['sequence'] = int(float(sequence))
                active = cell('active')
                if active is not None:
                    vals['active'] = active.lower() in (
                        'yes', 'true', '1', 'y', 'x', 'có')

                record = existing.get(code)
                if record:
                    record.write(vals)
                    sheet_updated += 1
                else:
                    record = Value.create(dict(
                        vals, code=code, category_id=category.id,
                        name=vals.get('name') or code))
                    sheet_created += 1
                    existing[code] = record

                name_vi = cell('name_vi')
                if name_vi:
                    # Writes the vi_VN translation through the alias mirror, so
                    # one sheet carries both languages.
                    record.name_vi = name_vi

            created += sheet_created
            updated += sheet_updated
            lines.append(_('%(category)s: %(created)s created, '
                           '%(updated)s updated.',
                           category=category.name,
                           created=sheet_created, updated=sheet_updated))

        summary = _('%(created)s value(s) created, %(updated)s updated, '
                    '%(skipped)s sheet(s) skipped.',
                    created=created, updated=updated, skipped=skipped)
        _logger.info('lookup import: %s', summary)
        self.write({'state': 'done', 'log': summary + '\n\n' + '\n'.join(lines)})
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
