#!/usr/bin/env python3
"""Migration audit report generator.

Reads Migration/Booking_mig.xlsx, Contact_mig.xlsx, Lookup.xlsx plus the
vietuat DB extracts in scratchpad/db/, replays the migration_runner transforms
cell-by-cell, and writes a colour-coded multi-sheet audit workbook.
"""
import csv, json, hashlib, re, sys, unicodedata
from datetime import datetime, date, time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = '/Users/adity/Documents/GitHub/health19'
DB = '/private/tmp/claude-502/-Users-adity-Documents-GitHub-health19/28f4b4cc-1b87-4ddc-90d3-83b619360cc5/scratchpad/db'
OUT = REPO + '/Migration/Migration_Audit_Report.xlsx'

# ------------------------------------------------------------------ styles
F_HDR = Font(bold=True, color='FFFFFF', size=10)
FILL_HDR = PatternFill('solid', fgColor='333F4F')
GREEN = PatternFill('solid', fgColor='C6EFCE'); F_GREEN = Font(color='006100', size=10)
YELLOW = PatternFill('solid', fgColor='FFEB9C'); F_YELLOW = Font(color='9C6500', size=10)
RED = PatternFill('solid', fgColor='FFC7CE'); F_RED = Font(color='9C0006', bold=True, size=10)
GREY = PatternFill('solid', fgColor='D9D9D9'); F_GREY = Font(color='595959', size=10)
BLUE = PatternFill('solid', fgColor='B4C6E7'); F_BLUE = Font(color='1F3864', bold=True, size=10)
ORANGE = PatternFill('solid', fgColor='F8CBAD'); F_ORANGE = Font(color='974706', bold=True, size=10)
F_PLAIN = Font(size=10)
F_TITLE = Font(bold=True, size=14)
F_SECT = Font(bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Border(*[Side(style='thin', color='BFBFBF')]*4)

STYLE = {'green': (GREEN, F_GREEN), 'yellow': (YELLOW, F_YELLOW), 'red': (RED, F_RED),
         'grey': (GREY, F_GREY), 'blue': (BLUE, F_BLUE), 'orange': (ORANGE, F_ORANGE)}

def paint(c, cat):
    if cat and cat in STYLE:
        c.fill, c.font = STYLE[cat]
    else:
        c.font = F_PLAIN

# ------------------------------------------------------------------ helpers (runner replicas)
def _norm(s):
    return (str(s).strip().lower() if s not in (None, False) else '')

def i18n(v):
    if isinstance(v, str) and v.startswith('{'):
        try:
            d = json.loads(v)
            return d.get('en_US') or next(iter(d.values()))
        except Exception:
            return v
    return v

def i18n_vi(v):
    if isinstance(v, str) and v.startswith('{'):
        try:
            d = json.loads(v)
            return d.get('vi_VN') or d.get('en_US') or next(iter(d.values()))
        except Exception:
            return v
    return v

GENDER_MAP = {'nữ': 'female', 'nu': 'female', 'nam': 'male', 'male': 'male',
              'female': 'female', 'không xác định': '', 'khong xac dinh': '',
              'khác': '', 'other': 'other'}
PAYMENT_METHOD_MAP = {'tiền mặt': 'cash', 'tien mat': 'cash', 'chuyển khoản': 'bank_transfer',
                      'chuyen khoan': 'bank_transfer', 'tiền mặt + chuyển khoản': 'other'}
CITY_MAP = {'hcmc': 'hcm', 'hcm': 'hcm', 'ho chi minh': 'hcm', 'hồ chí minh': 'hcm',
            'tp hcm': 'hcm', 'tphcm': 'hcm', 'hanoi': 'hn', 'hn': 'hn', 'hà nội': 'hn', 'ha noi': 'hn'}
FSO_STATUS_MAP = {'đã hoàn thành': 'completed / stage Completed', 'da hoan thanh': 'completed / stage Completed',
                  'huỷ': 'cancelled / stage Cancelled', 'hủy': 'cancelled / stage Cancelled', 'huy': 'cancelled / stage Cancelled',
                  'mới': 'confirmed / stage Booked', 'moi': 'confirmed / stage Booked',
                  'dời lịch': 'confirmed / stage Booked', 'doi lich': 'confirmed / stage Booked'}
LEAD_STATUS_MAP = {'mới': 'lead', 'moi': 'lead', 'đặt lịch hẹn': 'booking', 'dat lich hen': 'booking',
                   'huỷ': 'lost_booking', 'hủy': 'lost_booking'}
LEAD_SOURCE_MAP = {'hotline': 'phone_inquiry', 'zalo': 'zalo_marketing', 'khách hàng cũ': 'referral_patient'}
NON_STAFF = ('diag', 'greenlab', 'lab', 'diagnostic')

def norm_phone(value):
    if not value:
        return value
    digits = re.sub(r'\D', '', str(value))
    if digits.startswith('84') and len(digits) in (11, 12):
        rest = digits[2:]
        digits = rest if rest.startswith('0') else '0' + rest
    if digits.startswith('0'):
        if len(digits) == 10 and digits[1] != '0':
            return digits
    else:
        if len(digits) == 9:
            return '0' + digits
    return None  # invalid

def parse_date_any(v):
    if not v:
        return None
    s = str(v).strip()
    try:
        return datetime.fromisoformat(s).date()
    except Exception:
        pass
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    return None

def parse_time_hm(v):
    if not v:
        return None
    m = re.match(r'\s*(\d{1,2}):(\d{2})', str(v))
    return (int(m.group(1)), int(m.group(2))) if m else None

def normalize_nurse(seg):
    s = (seg or '').strip().replace('–', '-').replace('—', '-')
    s = re.sub(r'\s+', ' ', s)
    s = re.sub(r'^(HCMC|HCM|HANOI|HN)\s*-\s*', '', s, flags=re.I)
    is_doctor = bool(re.match(r'^\s*(BS|Dr)\b', s, flags=re.I))
    s = re.sub(r'^(ĐDPT|DDPT|ĐDT|DDT|ĐD|DD|DT|BSCK[I0-9]*|BS|Dr)[\s._]+', '', s, flags=re.I)
    s = re.sub(r'^part[\s_-]?time\s+', '', s, flags=re.I)
    s = re.sub(r'[\s_-]*oncall\s*$', '', s, flags=re.I)
    return s.strip(' _-'), is_doctor

def cell_to_json(v):
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime('%H:%M')
    return v

def booking_ref(rec):
    basis = '|'.join(str(rec.get(k) or '') for k in
                     ('Client ID', 'Appointment Date', 'time_in', 'dich_vu', 'CreatedOn'))
    return 'BKG-' + hashlib.md5(basis.encode('utf-8')).hexdigest()[:16]

# ------------------------------------------------------------------ load DB extracts
def load_csv(name):
    with open(f'{DB}/{name}.csv', newline='') as f:
        return list(csv.DictReader(f))

partners = load_csv('partners')
fsos = load_csv('fso')
leads = load_csv('leads')
xref = load_csv('xref')
service_types = load_csv('service_types')
staff = load_csv('staff')
products = load_csv('products')
sale_orders = load_csv('sale_orders')
payments = load_csv('payments')

db_client_codes = {r['legacy_client_code'] for r in partners}
db_refs = {r['legacy_booking_ref'] for r in fsos}
db_guids = {r['legacy_contact_guid'] for r in leads}
db_so_origins = {r['origin'] for r in sale_orders}
db_pay_refs = {r['legacy_booking_ref'] for r in payments}
xref_staff = {r['legacy_key'][6:]: r for r in xref if r['target_model'] == 'hr.employee'}
xref_svc = {r['legacy_key'][8:]: r for r in xref if r['target_model'] == 'health.service.type'}
xref_prod = {r['legacy_key'][8:]: r for r in xref if r['target_model'] == 'product.product'}
svc_names_vi = {_norm(i18n(r['name'])) for r in service_types}
prod_names = {_norm(i18n(r['name'])) for r in products}

# ------------------------------------------------------------------ load source workbooks
def load_sheet(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr_raw = [h if h is not None else 'col' for h in rows[0]]
    # disambiguate duplicates the same way the converter did
    seen, hdr = {}, []
    for h in hdr_raw:
        if h in seen:
            seen[h] += 1
            hdr.append('%s_%d' % (h, seen[h]))
        else:
            seen[h] = 0
            hdr.append(h)
    data = [r for r in rows[1:] if not all(c in (None, '') for c in r)]
    return hdr, data

bk_hdr, bk_rows = load_sheet(f'{REPO}/Migration/Booking_mig.xlsx')
ct_hdr, ct_rows = load_sheet(f'{REPO}/Migration/Contact_mig.xlsx')

def rec_of(hdr, row):
    return {hdr[i]: cell_to_json(row[i]) for i in range(len(hdr))}

# ------------------------------------------------------------------ column specs
# cat: key | asis | transform | special:<name> | notmig ; (target, transform note, extra note)
BK_SPEC = {
    'Sequence': ('notmig', '', 'Row number only', ''),
    'Appointment Date': ('special:appt_date', 'health.fieldservice.order.scheduled_datetime', 'Combined with time_in; VN local time converted to UTC', ''),
    'Appointment Time': ('notmig', '', 'Always 00:00 in the export; time_in holds the real time', ''),
    'CreatedOn': ('special:createdon', 'health.payment.transaction.transaction_date', 'Parsed dd/mm/yyyy HH:MM -> UTC; used as the payment date. FSO create date NOT preserved', ''),
    'dich_vu_tai': ('special:svcloc', 'health.fieldservice.order.service_location', "HARDCODED to 'home' for every row", "21 'Tại Phòng khám' rows migrated WRONG (system supports 'clinic')"),
    'CreatedBy': ('notmig', '', 'Legacy operator name not migrated', ''),
    'lien_ket_lien_he': ('notmig', '', 'Pancake link column (empty)', ''),
    'Contact_ad_id_tiktok': ('notmig', '', 'Ad tracking', ''),
    'Contact_dich_vu_danh_muc': ('notmig', '', 'Empty in export', ''),
    'Contact_dich_vu_quan_tam': ('notmig', '', 'Empty in export', ''),
    'so_cccd': ('asis', 'res.partner.national_id', 'Copied as-is', ''),
    'Name': ('notmig', '', "Short label ('Bn'); real name taken from ten_benh_nhan", ''),
    'so_dien_thoai': ('special:phone', 'res.partner.mobile', 'normalize_vn_phone; export truncated phones to 3-4 digits so ALL failed -> random placeholder numbers were set instead', ''),
    'so_dien_thoai_2': ('notmig', '', 'Second phone not migrated', ''),
    'gui_xe': ('notmig', '', 'Parking fee not migrated', ''),
    'Contact_nam_sinh': ('special:dob', 'res.partner.birth_date', "Parsed 'HH:MM dd/mm/yyyy' -> date (time prefix dropped)", ''),
    'nghe_nghiep': ('asis', 'res.partner.profession', 'Copied as-is', ''),
    'dan_toc': ('notmig', '', 'Ethnicity: no target field yet', ''),
    'Contact_gioi_tinh': ('special:gender', 'res.partner.gender', "Mapped Nữ->female, Nam->male; 'Không xác định'/'Khác' DROPPED to blank", ''),
    'Client ID': ('special:clientid', 'res.partner.legacy_client_code', 'Upsert key (verbatim)', ''),
    'ten_benh_nhan': ('special:pname', 'res.partner.name', "Copied; blank -> fallback 'BN <code>'", ''),
    'SourceofClient': ('notmig', '(res.partner / booking_source)', "NOT migrated - booking_source hardcoded 'phone'. Real acquisition source dropped; backfill recommended", ''),
    'Client_Type': ('notmig', '', "New/Repeat flag not migrated (no target field)", ''),
    'lien_ket_benh_nhan': ('notmig', '', 'Pancake link column (empty)', ''),
    'trang_thai': ('special:fsostatus', 'health.fieldservice.order.state + stage_id', "Closed map; 'Dời lịch' collapsed into Confirmed (reschedule flag lost)", ''),
    'ly_do_huy': ('special:cancel', 'health.fieldservice.order.cancellation_notes', 'Free text only on cancelled rows; NOT linked to health.booking.cancellation.reason lookup', ''),
    'Contact_phong_kham': ('notmig', '', 'City name variant; phong_kham_city used instead', ''),
    'phong_kham_city': ('special:city', 'res.partner.catchment_province_id + primary_facility_id / fso.facility_id', 'HCMC->TPHCM(02), Hanoi->Hà Nội(01); facility resolved per city', ''),
    'phu_phi': ('notmig', '', 'SURCHARGE AMOUNT NOT MIGRATED (financial)', ''),
    'nguoi_thuc_hien': ('notmig', '', 'Performer not migrated (bac_si_dieu_duong used for staff)', ''),
    'khoang_cach': ('notmig', '', 'Distance re-computed by geocoding instead (clinic_drive_distance_km)', ''),
    'Contact_FullAddress': ('special:fulladdr', 'res.partner.street (fallback)', 'Used only when Contact_Address is empty', ''),
    'country': ('notmig', '', 'Not migrated (address kept as one street line)', ''),
    'province': ('notmig', '', 'Not migrated', ''),
    'district': ('notmig', '', 'Not migrated', ''),
    'commune': ('notmig', '', 'Not migrated', ''),
    'Contact_Address': ('special:addr', 'res.partner.street', 'Copied as-is', ''),
    'note': ('special:note', 'health.fieldservice.order.patient_notes', 'Concatenated with mo_ta_tinh_trang', ''),
    'Contact_customer_id': ('notmig', '', 'Pancake id', ''),
    'phong_kham_Owner': ('notmig', '', 'Legacy owner', ''),
    'loai_kham': ('notmig', '', 'Empty in export', ''),
    'dich_vu': ('special:services', 'fso.appointment_type_id + sale.order.line + health.service.type / product.product (auto-created)', "';'-split. First service -> appointment type; every service -> sale line at PRICE 0. 68 service types + 166 products auto-created with 0 price", ''),
    'time_in': ('special:timein', 'health.fieldservice.order.scheduled_datetime (time part)', 'Combined with Appointment Date', ''),
    'time_out': ('special:timeout', 'health.fieldservice.order.scheduled_duration', 'time_out - time_in in minutes; fallback 60 when missing/invalid', ''),
    'TransferMoney': ('special:money', 'health.payment.transaction.amount (summed with Cash)', 'One combined payment record; split by method lost when both used', ''),
    'thanh_toan_truoc': ('notmig', '', 'PREPAYMENT AMOUNT NOT MIGRATED (financial)', ''),
    'Cash': ('special:money', 'health.payment.transaction.amount (summed with TransferMoney)', '', ''),
    'SubTotal': ('notmig', '', 'NOT MIGRATED - sale lines were created at price 0 (revenue amounts lost)', ''),
    'Discount': ('notmig', '', 'NOT MIGRATED (financial)', ''),
    'thanh_toan': ('notmig', '', "Payment-status text not migrated; transaction status hardcoded 'collected'", ''),
    'ModifiedBy': ('notmig', '', '', ''),
    'thoi_gian_bat_dau': ('notmig', '', 'First-use date not migrated', ''),
    'bac_si_dieu_duong': ('special:staff', 'fso.assigned_staff_ids -> hr.employee (auto-stub)', "';'-split; city prefix + role tokens (ĐD/BS/...) stripped; lab tokens (GREENLAB/DIAG) skipped; 48 staff stubs auto-created", ''),
    'tao_don': ('notmig', '', "Always 'Tạo đơn'", ''),
    'so_don_da_tao': ('notmig', '', 'Order count implicit (1 SO per booking)', ''),
    'dich_vu_don_vi_tinh': ('notmig', '', 'Unit of measure not migrated (see Lookup sheet)', ''),
    'GrandTotal': ('notmig', '', 'NOT MIGRATED (financial; 0 in export)', ''),
    'TotalQuantity': ('notmig', '', 'NOT MIGRATED - every sale line qty hardcoded 1', ''),
    'hinh_thuc_thanh_toan': ('special:paymethod', 'health.payment.transaction.payment_method', "Tiền mặt->cash, Chuyển khoản->bank_transfer, mixed->'other'", ''),
    'tra': ('notmig', '', 'Paid flag not migrated (financial)', ''),
    'kanban_order': ('notmig', '', '', ''),
    'nguoi_phu_trach': ('notmig', '', 'Person in charge not migrated', ''),
    'ModifiedOn': ('notmig', '', '', ''),
    'in_don': ('notmig', '', 'Invoice-print flag: no target concept (empty in export)', ''),
    'Owner': ('notmig', '', '', ''),
    'PancakeCustomer ID': ('notmig', '', 'Pancake id', ''),
    'page_id': ('notmig', '', 'Channel/ad tracking', ''),
    'ad_id': ('notmig', '', 'Channel/ad tracking', ''),
    'TimeEnd': ('notmig', '', '', ''),
    'TimeStart': ('notmig', '', '', ''),
    'mo_ta_tinh_trang': ('special:note', 'health.fieldservice.order.patient_notes', 'Concatenated with note', ''),
    'ad_id_fb': ('notmig', '', 'Channel/ad tracking', ''),
    'conversation_id': ('notmig', '', 'Channel/ad tracking', ''),
    'Facebook Link': ('notmig', '', 'Channel/ad tracking', ''),
}

CT_SPEC = {
    'STT': ('notmig', '', 'Row number only', ''),
    'ID': ('special:guid', 'crm.lead.legacy_contact_guid', 'Upsert key (verbatim)', ''),
    'ngay_dau_su_dung': ('notmig', '', 'First-use date not migrated', ''),
    'Name': ('special:lname', 'crm.lead.name', "Copied; blank -> fallback 'Contact <guid8>'", ''),
    'gioi_tinh': ('special:lgender', '(none)', 'NOT MIGRATED - crm.lead has no gender field (value validated then dropped)', ''),
    'nam_sinh': ('notmig', '', 'No DOB field on crm.lead', ''),
    'nguon_khach_hang': ('special:lsource', 'crm.lead.healthcare_lead_source', "Only Hotline/Zalo/'Khách hàng cũ' mapped; every other source DROPPED (field left empty)", ''),
    'phong_kham': ('notmig', '', 'Clinic/branch not stored on the lead', ''),
    'customer_id': ('notmig', '', 'Pancake id', ''),
    'ly_do_tu_choi': ('special:reject', 'crm.lead.reason_if_rejected', 'Copied as free text; NOT linked to a structured lost-reason lookup', ''),
    'Phone': ('special:phone', 'crm.lead.phone', 'normalize_vn_phone; truncated numbers failed -> random placeholder numbers were set instead', ''),
    'so_dien_thoai_2': ('notmig', '', 'Second phone not migrated', ''),
    'NextContactAt': ('notmig', '', 'Not migrated (no follow-up activity created)', ''),
    'so_cccd': ('notmig', '', 'No national-ID field on crm.lead', ''),
    'ten_benh_nhan': ('notmig', '', 'Patient name not migrated (lead is pre-patient)', ''),
    'khoang_cach': ('notmig', '', 'Distance not migrated', ''),
    'dich_vu_quan_tam': ('special:desc', 'crm.lead.description (concat)', 'Appended to description; structured service_interest selection NOT set', ''),
    'loai_khach_hnagf': ('notmig', '', 'New/Repeat flag not migrated', ''),
    'ghi_chu': ('special:desc', 'crm.lead.description (concat)', '', ''),
    'Source': ('notmig', '', "Pancake inbox tag ('INBOX') not migrated", ''),
    'cskh_3': ('notmig', '', 'Care-call disposition not migrated (health_contact_outcome not set)', ''),
    'Untitled': ('notmig', '', 'Unnamed legacy column', ''),
    'Untitled2': ('notmig', '', 'Unnamed legacy column', ''),
    'phan_hoi_kh': ('notmig', '', 'Customer feedback not migrated', ''),
    'PancakeCustomer': ('notmig', '', 'Pancake id', ''),
    'PancakeCustomer ID': ('notmig', '', 'Pancake id', ''),
    'Address': ('notmig', '', 'NOT migrated - import_contacts never writes an address', ''),
    'Status': ('special:lstatus', 'crm.lead.contact_status + stage_id', "Mới->Lead, Đặt lịch hẹn->Booking, Huỷ->Lost Booking; anything else would fall back to 'active'", ''),
    'Owner': ('notmig', '', 'Owner/salesperson not migrated', ''),
    'pancake_tag': ('notmig', '', 'Pancake tags not migrated', ''),
    'lien_ket_lich_hen': ('notmig', '', 'Link column', ''),
    'FullAddress': ('notmig', '', 'Not migrated', ''),
    'Address_1': ('notmig', '', 'Duplicate address column; not migrated', ''),
    'country': ('notmig', '', '', ''), 'province': ('notmig', '', '', ''),
    'district': ('notmig', '', '', ''), 'commune': ('notmig', '', '', ''),
    'kanban_order': ('notmig', '', '', ''),
    'dich_vu_don_vi_tinh': ('notmig', '', 'Unit of measure (see Lookup sheet)', ''),
    'dich_vu_danh_muc': ('notmig', '', 'Service category (see Lookup sheet)', ''),
    'CreatedOn': ('notmig', '', 'Legacy create date not preserved', ''),
    'dich_vu_RetailPrice': ('notmig', '', 'NOT MIGRATED (financial)', ''),
    'dich_vu': ('notmig', '', 'Services on the contact not migrated (bookings carry the services)', ''),
    'SubTotal': ('notmig', '', 'NOT MIGRATED (financial)', ''),
    'Discount': ('notmig', '', 'NOT MIGRATED (financial)', ''),
    'GrandTotal': ('notmig', '', 'NOT MIGRATED (financial)', ''),
    'CreatedBy': ('notmig', '', '', ''),
    'ModifiedOn': ('notmig', '', '', ''),
    'psid': ('notmig', '', 'Channel/ad tracking', ''),
    'TotalQuantity': ('notmig', '', 'NOT MIGRATED (financial)', ''),
    'Email': ('special:email', 'crm.lead.email_from', 'NOT written by the importer (all values empty in this export)', ''),
    'ModifiedBy': ('notmig', '', '', ''),
    'page_id': ('notmig', '', 'Channel/ad tracking', ''),
    'LastContactUser': ('notmig', '', '', ''),
    'ad_id': ('notmig', '', 'Channel/ad tracking', ''),
    'TransferMoney': ('notmig', '', 'NOT MIGRATED (financial; contact-level)', ''),
    'Note': ('special:desc', 'crm.lead.description (concat)', '', ''),
    'Cash': ('notmig', '', 'NOT MIGRATED (financial; contact-level)', ''),
    'LastContactAt': ('notmig', '', '', ''),
    'ad_id_tiktok': ('notmig', '', 'Channel/ad tracking', ''),
    'ad_id_fb': ('notmig', '', 'Channel/ad tracking', ''),
    'Facebook Link': ('notmig', '', 'Channel/ad tracking', ''),
}

# ------------------------------------------------------------------ cell classifiers
bk_issues, ct_issues = [], []

def classify_booking_cell(col, val, rec, ref_ok):
    """Return (cat, issue_note or None). val is the json-converted value."""
    blank = val in (None, '')
    spec = BK_SPEC.get(col, ('notmig', '', '', ''))
    kind = spec[0]
    if kind == 'notmig':
        return 'grey', None
    if kind == 'asis':
        return (None if blank else 'green'), None
    k = kind.split(':')[1]
    key = str(rec.get('Client ID') or '?')
    if k == 'clientid':
        if blank:
            return 'red', 'Blank Client ID - entire row skipped by the importer'
        return ('blue', None) if val in db_client_codes else ('red', 'Client code NOT found in vietuat')
    if k == 'pname':
        if blank:
            return 'red', "Blank patient name - fallback 'BN %s' used" % key
        return 'green', None
    if k == 'phone':
        if blank:
            return None, None
        return 'red', "Phone '%s' truncated in export (invalid) - system now holds a RANDOM PLACEHOLDER number" % val
    if k == 'dob':
        if blank:
            return None, None
        d = parse_date_any(val)
        if d:
            return 'yellow', None
        return 'red', "Unparsable date of birth '%s' - birth_date left empty" % val
    if k == 'gender':
        if blank:
            return None, None
        g = GENDER_MAP.get(_norm(val), '__miss__')
        if g == '__miss__':
            return 'red', "Gender '%s' not in map - dropped" % val
        if g == '':
            return 'red', "Gender '%s' has no system equivalent - dropped to blank" % val
        return 'yellow', None
    if k == 'city':
        if blank:
            return 'red', 'Blank city - client would be skipped (no catchment)'
        return ('yellow', None) if _norm(val) in CITY_MAP else ('red', "City '%s' unmapped - client skipped" % val)
    if k == 'fsostatus':
        if blank:
            return 'red', 'Blank status - fell back to Draft'
        mapped = FSO_STATUS_MAP.get(_norm(val))
        if not mapped:
            return 'red', "Status '%s' unmapped - fell back to Draft" % val
        if _norm(val) in ('dời lịch', 'doi lich'):
            return 'yellow', "'Dời lịch' (rescheduled) collapsed into Confirmed - reschedule history lost"
        return 'yellow', None
    if k == 'cancel':
        if blank:
            return None, None
        if _norm(rec.get('trang_thai')) in ('huỷ', 'hủy', 'huy'):
            return 'yellow', "Cancel reason kept as free text only (not linked to cancellation-reason lookup)"
        return 'red', "Cancel reason '%s' on a NON-cancelled row - dropped" % val
    if k == 'svcloc':
        if blank:
            return None, None
        if _norm(val) in ('tại nhà', 'tai nha'):
            return 'yellow', None
        return 'red', "Location '%s' forced to 'home' - system supports 'clinic' but importer hardcoded home" % val
    if k == 'services':
        if blank:
            return None, None
        names = [s.strip() for s in str(val).split(';') if s.strip()]
        def svc_known(n):
            nn = _norm(n)
            return (nn[:120] in xref_svc or nn[:120] in xref_prod
                    or nn in svc_names_vi or nn in prod_names)
        missing = [n for n in names if not svc_known(n)]
        if missing:
            return 'red', 'Service(s) not found in vietuat: %s' % '; '.join(missing[:3])
        return 'yellow', None
    if k == 'staff':
        if blank:
            return None, None
        segs = [s.strip() for s in str(val).replace('–', '-').split(';') if s.strip()]
        kept, skipped = [], []
        for s in segs:
            nm, _doc = normalize_nurse(s)
            (skipped if (not nm or _norm(nm) in NON_STAFF) else kept).append(nm or s)
        if not kept:
            return 'red', "No real staff in '%s' (lab/diagnostic tokens) - no assignment created" % val
        unknown = [nm for nm in kept if _norm(nm) not in xref_staff]
        if unknown:
            return 'red', 'Staff not found in vietuat xref: %s' % '; '.join(unknown[:3])
        if skipped:
            return 'yellow', 'Lab token(s) skipped: %s' % '; '.join(skipped[:2])
        return 'yellow', None
    if k in ('timein', 'timeout'):
        if blank:
            return ('red', 'Missing %s - duration fell back to 60 min' % col) if k == 'timeout' and rec.get('time_in') else (None, None)
        t = parse_time_hm(val)
        if not t:
            return 'red', "Unparsable time '%s'" % val
        if k == 'timeout':
            a = parse_time_hm(rec.get('time_in'))
            if a and ((t[0]*60+t[1]) - (a[0]*60+a[1])) <= 0:
                return 'red', 'time_out <= time_in - duration fell back to 60 min'
        return 'yellow', None
    if k == 'appt_date':
        if blank:
            return 'red', 'Missing appointment date - scheduled_datetime left empty'
        return 'yellow', None
    if k == 'createdon':
        return (None, None) if blank else ('yellow', None)
    if k == 'money':
        if blank:
            return None, None
        try:
            amt = float(str(val).replace(',', '').strip() or 0)
        except Exception:
            return 'red', "Unparsable amount '%s' - treated as 0" % val
        return ('yellow', None) if amt > 0 else (None, None)
    if k == 'paymethod':
        if blank:
            return None, None
        m = PAYMENT_METHOD_MAP.get(_norm(val))
        if m == 'other':
            return 'red', "Mixed method '%s' collapsed to 'other' (cash/bank split lost)" % val
        if not m:
            return 'red', "Payment method '%s' unmapped - fell back to 'other'" % val
        return 'yellow', None
    if k == 'fulladdr':
        if blank:
            return None, None
        return ('yellow', 'Used as street fallback') if not rec.get('Contact_Address') else ('grey', None)
    if k in ('addr',):
        return (None if blank else 'green'), None
    if k == 'note':
        return (None if blank else 'yellow'), None
    return None, None

def classify_contact_cell(col, val, rec):
    blank = val in (None, '')
    spec = CT_SPEC.get(col, ('notmig', '', '', ''))
    kind = spec[0]
    if kind == 'notmig':
        return 'grey', None
    if kind == 'asis':
        return (None if blank else 'green'), None
    k = kind.split(':')[1]
    if k == 'guid':
        if blank:
            return 'red', 'Blank GUID - row skipped by the importer'
        return ('blue', None) if val in db_guids else ('red', 'GUID not found in vietuat')
    if k == 'lname':
        if blank:
            return 'red', "Blank name - fallback 'Contact <guid>' used"
        return 'green', None
    if k == 'phone':
        if blank:
            return None, None
        return 'red', "Phone '%s' truncated/invalid in export - system now holds a RANDOM PLACEHOLDER number" % val
    if k == 'lgender':
        if blank:
            return None, None
        return 'red', "Gender '%s' NOT migrated - crm.lead has no gender field" % val
    if k == 'lsource':
        if blank:
            return None, None
        if _norm(val) in LEAD_SOURCE_MAP:
            return 'yellow', None
        return 'red', "Lead source '%s' unmapped - healthcare_lead_source left EMPTY" % val
    if k == 'lstatus':
        if blank:
            return 'red', "Blank status - fell back to 'Initial Contact'"
        if _norm(val) in LEAD_STATUS_MAP:
            return 'yellow', None
        return 'red', "Status '%s' unmapped - fell back to 'Initial Contact'" % val
    if k == 'reject':
        return (None if blank else 'yellow'), ('Kept as free text; not linked to a structured lookup' if not blank else None)
    if k == 'desc':
        return (None if blank else 'yellow'), None
    if k == 'email':
        if blank:
            return None, None
        return 'red', 'Email not written by the importer'
    return None, None

# ------------------------------------------------------------------ build workbook
out = openpyxl.Workbook()
out.remove(out.active)

def style_header(ws, row=1):
    for c in ws[row]:
        if c.value not in (None, ''):
            c.font = F_HDR; c.fill = FILL_HDR; c.alignment = Alignment(wrap_text=True, vertical='center')

# ---------- data mirror sheets
def build_mirror(name, hdr, rows, spec, classify, issues, keyfn):
    ws = out.create_sheet(name)
    ws.append(hdr)
    style_header(ws)
    ws.freeze_panes = 'A2'
    counts = {}
    for ridx, row in enumerate(rows, start=2):
        rec = rec_of(hdr, row)
        for cidx, col in enumerate(hdr, start=1):
            val = cell_to_json(row[cidx-1]) if cidx-1 < len(row) else None
            c = ws.cell(row=ridx, column=cidx, value=row[cidx-1] if cidx-1 < len(row) else None)
            if classify is classify_booking_cell:
                cat, note = classify(col, val, rec, True)
            else:
                cat, note = classify(col, val, rec)
            paint(c, cat)
            if cat:
                counts[cat] = counts.get(cat, 0) + 1
            if note:
                issues.append((ridx, keyfn(rec), col, '' if val is None else str(val)[:80], cat, note))
    for cidx in range(1, len(hdr)+1):
        ws.column_dimensions[get_column_letter(cidx)].width = 14
    return counts

bk_counts = build_mirror('Booking Data', bk_hdr, bk_rows, BK_SPEC, classify_booking_cell,
                         bk_issues, lambda r: str(r.get('Client ID') or ''))
ct_counts = build_mirror('Contact Data', ct_hdr, ct_rows, CT_SPEC, classify_contact_cell,
                         ct_issues, lambda r: str(r.get('ID') or '')[:12])

# ---------- DB cross-check per row (booking refs / sale orders / payments)
missing_refs, no_so, no_pay = [], [], []
for i, row in enumerate(bk_rows, start=2):
    rec = rec_of(bk_hdr, row)
    ref = booking_ref(rec)
    if ref not in db_refs:
        missing_refs.append((i, ref))
        bk_issues.append((i, str(rec.get('Client ID') or ''), '(row)', ref, 'red',
                          'Booking NOT found in vietuat (legacy ref missing)'))
    else:
        if str(rec.get('dich_vu') or '').strip() and ref not in db_so_origins:
            no_so.append((i, ref))
            bk_issues.append((i, str(rec.get('Client ID') or ''), 'dich_vu', str(rec.get('dich_vu'))[:60],
                              'red', 'No sale order found in vietuat for this booking'))
        cash = rec.get('Cash'); xfer = rec.get('TransferMoney')
        def fnum(v):
            try: return float(str(v).replace(',', '').strip() or 0)
            except Exception: return 0.0
        if fnum(cash) + fnum(xfer) > 0 and ref not in db_pay_refs:
            no_pay.append((i, ref))
            bk_issues.append((i, str(rec.get('Client ID') or ''), 'Cash/TransferMoney',
                              '%s / %s' % (cash, xfer), 'red', 'No payment transaction found in vietuat'))
missing_guids = []
for i, row in enumerate(ct_rows, start=2):
    rec = rec_of(ct_hdr, row)
    g = str(rec.get('ID') or '').strip()
    if g and g not in db_guids:
        missing_guids.append((i, g))

# ---------- mapping sheets
def build_mapping(name, hdr, rows, spec_map):
    ws = out.create_sheet(name)
    ws.append(['Legacy Column', 'Sample Value', 'Migrated?', 'Target Model.Field', 'Transform / Notes'])
    style_header(ws)
    ws.freeze_panes = 'A2'
    # sample: first non-blank value
    for col in hdr:
        idx = hdr.index(col)
        sample = ''
        for r in rows:
            v = r[idx] if idx < len(r) else None
            if v not in (None, ''):
                sample = str(cell_to_json(v))[:40]
                break
        kind, target, note, extra = spec_map.get(col, ('notmig', '', 'Unknown column', ''))
        if kind == 'notmig':
            status, cat = 'NO - not migrated', 'grey'
        elif kind == 'asis':
            status, cat = 'YES - as-is', 'green'
        elif kind.startswith('special:phone'):
            status, cat = 'LOST - placeholder set', 'red'
        elif kind in ('special:lgender', 'special:email'):
            status, cat = 'NO - dropped', 'red'
        elif kind in ('special:clientid', 'special:guid'):
            status, cat = 'YES - upsert key', 'blue'
        else:
            status, cat = 'YES - transformed', 'yellow'
        full_note = (note + ('. ' + extra if extra else '')).strip()
        ws.append([col, sample, status, target, full_note])
        for c in ws[ws.max_row]:
            c.alignment = WRAP; c.border = THIN
        paint(ws.cell(row=ws.max_row, column=3), cat)
    widths = [24, 28, 22, 42, 70]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws

build_mapping('Mapping - Booking', bk_hdr, bk_rows, BK_SPEC)
build_mapping('Mapping - Contact', ct_hdr, ct_rows, CT_SPEC)

# ---------- issues sheets
def build_issues(name, issues):
    ws = out.create_sheet(name)
    ws.append(['Source Row', 'Key', 'Column', 'Legacy Value', 'Severity', 'What happened'])
    style_header(ws)
    ws.freeze_panes = 'A2'
    sev = {'red': 'LOST / FAILED', 'yellow': 'CHANGED', 'blue': 'KEY', None: ''}
    for (r, key, col, val, cat, note) in issues:
        ws.append([r, key, col, val, sev.get(cat, cat or ''), note])
        paint(ws.cell(row=ws.max_row, column=5), cat)
        for c in ws[ws.max_row]:
            c.alignment = WRAP
    widths = [10, 14, 20, 34, 16, 80]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

build_issues('Booking Issues', bk_issues)
build_issues('Contact Issues', ct_issues)

# ---------- lookup analysis sheets ---------------------------------------
uom = load_csv('uom')
cancel_reasons = load_csv('health_booking_cancellation_reason')
utm = load_csv('utm_source')
catchments = load_csv('catchments')
facilities = load_csv('facilities')
lead_reasons = load_csv('health_lead_reason')
lost_reasons = load_csv('crm_lost_reason')
specialties = load_csv('specialties')

def LK(rows_):
    """rows_: list of (list_name, vi, en, target, exists, match, action)"""
    return rows_

E = 'EXISTS'; C = 'CREATE'; D = 'DECISION'
# --- Contact Table lookups
ct_lookup = []
def add(listname, vi, en, target, status, match, action=''):
    ct_lookup.append((listname, vi, en, target, status, match, action))

g = 'Giới tính BN (Gender)'
add(g, 'Nam', 'Male', 'res.partner.gender (selection)', E, "'male'")
add(g, 'Nữ', 'Female', 'res.partner.gender (selection)', E, "'female'")
add(g, 'Khác', 'Other', 'res.partner.gender (selection)', E, "'other'", "Exists BUT the importer dropped it to blank - remap")
add(g, 'Không xác định', 'Unspecified', 'res.partner.gender (selection)', D, "nearest: 'prefer_not_to_say'",
    'No exact equivalent - decide: map to prefer_not_to_say or add a selection value (code change)')

r = 'Lý do từ chối (Rejection Reason)'
rej = [('Spam', 'Spam', E, "contact_status 'Spam Call' exists on crm.lead", 'Map to status, not a reason record'),
       ('Không có bác sĩ <2.5h', 'No doctor available < 2.5 hours', C, '', ''),
       ('Không có bác sĩ >2.5h', 'No doctor available > 2.5 hours', C, '', ''),
       ('Không có điều dưỡng <2.5h', 'No nurse available < 2.5 hours', C, '', ''),
       ('Không có điều dưỡng >2.5h', 'No nurse available > 2.5 hours', C, '', ''),
       ('Giá quá cao', 'Price too high', E, "crm.lost.reason 'Too expensive / Quá đắt'", ''),
       ('Vẫn đang cân nhắc', 'Still considering', C, '', ''),
       ('Không có đơn thuốc', 'No prescription', C, '', ''),
       ('Chỉ cần tư vấn', 'Consultation only', C, '', ''),
       ('Địa chỉ quá xa', 'Address too far', C, '', ''),
       ('Không có dịch vụ', 'Service not available', C, '', ''),
       ('Nhầm số hotline', 'Wrong hotline number', C, '', ''),
       ('Không có máy móc thiết bị', 'No equipment available', C, '', ''),
       ('Không nói', 'No response / Did not say', C, '', ''),
       ('Khác', 'Other', C, '', '')]
for vi, en, st, match, act in rej:
    add(r, vi, en, 'crm.lost.reason (lost-reason lookup); today stored as free text in crm.lead.reason_if_rejected',
        st, match, act or ('Create as crm.lost.reason (or health.lead.reason) and link' if st == C else ''))

s = 'Dịch vụ quan tâm (Service of Interest)'
svc_int = [('Chăm sóc giảm nhẹ', 'Palliative care', E, "service_interest 'palliative'"),
           ('Chăm sóc cá nhân', 'Personal care', C, ''),
           ('Chăm sóc vết thương', 'Wound care', C, ''),
           ('Bác sĩ khám', 'Medical Examination', E, "service_interest 'consultation' (approx)"),
           ('Tiêm', 'Injection', C, ''),
           ('Truyền', 'IV infusion', C, ''),
           ('Thụt tháo', 'Enema', C, ''),
           ('Đặt, rút sonde', 'Catheter insertion/removal', C, ''),
           ('Hút đờm, vỗ rung đờm', 'Sputum suction, chest percussion', C, ''),
           ('Xét nghiệm', 'Laboratory testing', C, ''),
           ('Khác', 'Other', C, '')]
for vi, en, st, match in svc_int:
    add(s, vi, en, 'crm.lead.service_interest (selection, 8 generic values)', st, match,
        '' if st == E else 'No matching selection value - add selection values (code) or model as health.service.type categories')

t = 'Loại khách hàng (Customer Type)'
add(t, 'Mới', 'New', '(no target field)', D, '', 'Client_Type was not migrated; new/repeat is derivable from booking history - decide if a field is needed')
add(t, 'Cũ', 'Repeat', '(no target field)', D, '', 'Same as above')

c1 = 'CSKH / Contact Disposition'
cskh = [('Chốt dùng dịch vụ', 'Service confirmed', E, "health_contact_outcome 'service_booked'"),
        ('Không nghe máy', 'No answer', D, "nearest: 'no_response'"),
        ('Không phản hồi', 'No response', E, "health_contact_outcome 'no_response'"),
        ('Bận gọi lại sau', 'Busy - call back later', E, "health_contact_outcome 'pending_follow_up'"),
        ('Cân nhắc thêm', 'Considering further', E, "health_contact_outcome 'future_opportunity' (approx)"),
        ('Tham khảo dịch vụ', 'Service inquiry / reference only', D, "nearest: 'not_qualified'"),
        ('Từ chối dịch vụ', 'Service declined', E, "health_contact_outcome 'rejected'")]
for vi, en, st, match in cskh:
    add(c1, vi, en, 'crm.lead.health_contact_outcome (selection)', st, match,
        '' if st == E else 'Distinct concept - decide mapping or add selection value')

ss = 'Service Status (Contact)'
sst = [('Mới', 'New', E, "contact_status 'lead' (used by importer)"),
       ('Đặt lịch hẹn', 'Appointment scheduled', E, "contact_status 'booking'"),
       ('Đang suy nghĩ', 'Thinking / considering', C, ''),
       ('Liên hệ lại', 'To be contacted again', C, ''),
       ('Hủy', 'Cancelled', E, "contact_status 'lost_booking'"),
       ('Đã sử dụng', 'Service used', C, ''),
       ('Cũ', 'Existing', C, '')]
for vi, en, st, match in sst:
    add(ss, vi, en, 'crm.lead.contact_status (selection: active/booking/lead/lost_booking/spam)', st, match,
        '' if st == E else 'No equivalent status - would fall back to Initial Contact; add selection value if needed')

u = 'Đơn vị tính (Unit of Measure)'
uom_names = {_norm(i18n(x['name'])) for x in uom}
for vi, en, hint in [('Ca', 'Shift', None), ('Giờ', 'Hour', 'hours'), ('Lần', 'Visit / per occurrence', 'units'),
                     ('Km', 'Kilometer', 'km'), ('Chai', 'Bottle', None)]:
    if hint and hint in uom_names:
        add(u, vi, en, 'uom.uom', E, "uom '%s'" % hint.title())
    else:
        add(u, vi, en, 'uom.uom', C, '', 'Create UoM record')

n = 'Nguồn Khách Hàng (Lead Source)'
utm_names = {_norm(x['name']) for x in utm}
src = [('Form', 'Online form', E, "healthcare_lead_source 'website_form'", ''),
       ('Hotline', 'Hotline / Call-in', E, "healthcare_lead_source 'phone_inquiry' (used by importer)", ''),
       ('Đối tác', 'Partner', C, '', 'No partner source; create utm.source or add selection value'),
       ('Tự đến PK', 'Walk-in to clinic', E, "healthcare_lead_source 'walk_in'", 'Importer did NOT map it - remap'),
       ('Người Giới Thiệu', 'Referrer', E, "healthcare_lead_source 'referral_patient'", 'Importer did NOT map it - remap'),
       ('Facebook', 'Facebook', E, "healthcare_lead_source 'facebook_ad' + utm.source 'Facebook'", 'Importer did NOT map it - remap'),
       ('Zalo', 'Zalo', E, "healthcare_lead_source 'zalo_marketing' + utm.source 'Zalo' (used by importer)", ''),
       ('TikTok', 'TikTok', C, '', 'No TikTok source anywhere - create utm.source'),
       ('Google', 'Google', E, "utm.source 'Google Ads' / 'Google Organic'", 'No healthcare_lead_source value - decide'),
       ('Phòng khám Gia đình Việt Úc', '(sub-source: VAFC clinic)', C, '', '166 booking + 61 contact rows carry this - create utm.source'),
       ('Chăm sóc tại nhà Việt Úc', '(sub-source: home care)', C, '', 'Create utm.source'),
       ('Hanoi / HCMC', '(sub-source: city)', D, '', 'City is a catchment, not a source - decide'),
       ('FWD / Ivie - Bác sĩ ơi / Daiichi', '(partner sub-sources)', C, '', 'Create utm.source or partner referrer records'),
       ('Trường Quốc Tế Úc ACG / Trường Quốc Tế Châu Âu', '(school partners)', C, '', 'Create utm.source or partner referrer records'),
       ('Khách hàng cũ', 'Former client', E, "healthcare_lead_source 'referral_patient' (used by importer)", ''),
       ('Ms. Đặng Thị Thương / ĐD Song Phú / Bác sĩ/Điều dưỡng / ĐD Trúc Mai / ĐD Lê Đặng Anh Thương / Ms.Thu', '(named referrers)', D, '', 'Individual referrers - decide: utm.source records vs referral partner links')]
for vi, en, st, match, act in src:
    add(n, vi, en, 'crm.lead.healthcare_lead_source (selection) / utm.source (records)', st, match, act)

dc = 'Contact_dich_vu_danh_muc (Service Category)'
add(dc, 'DỊCH VỤ ĐIỀU DƯỠNG', 'Nursing services', 'health.service.type.category (selection)', E, "'nursing_care'")
add(dc, 'DỊCH VỤ BÁC SĨ', 'Doctor services', 'health.service.type.category (selection)', E, "'consultation' (approx)")
add(dc, 'DỊCH VỤ BÁC SĨ/ Người nước ngoài', 'Doctor services / foreigners', 'health.service.type.category (selection)', D, '',
    'Foreigner-doctor category has no equivalent - decide (pricing tier?)')

lk = 'loai_kham (Visit Type)'
add(lk, 'Khám nội', 'Internal medicine visit', 'health.medical.specialty', E, "'Internal Medicine / Nội bộ Y học'")
add(lk, 'Khám ngoại', 'Surgery/external visit', 'health.medical.specialty', C, '', "Create 'Surgery / Ngoại khoa' specialty")

# --- Booking Table lookups
bk_lookup = []
def addb(listname, vi, en, target, status, match, action=''):
    bk_lookup.append((listname, vi, en, target, status, match, action))

st_ = 'Trạng thái lịch hẹn (Booking Status)'
bst = [('Đặt chỗ mới', 'New Booking', E, "state 'confirmed' + stage 'Booked'", "Importer maps 'Mới'"),
       ('Lịch tiếp theo', 'Next appointment', D, '', 'No equivalent state - would fall back to Draft; decide mapping'),
       ('Đang tiến hành', 'In Progress', E, "state 'in_progress' EXISTS", "But NOT in the importer's map - would fall back to Draft; add to map"),
       ('Đã hoàn thành', 'Completed', E, "state 'completed' + stage 'Completed' (used: 911 rows)", ''),
       ('Dời lịch', 'Rescheduled', D, "collapsed into 'confirmed'", 'System has no Rescheduled state - 5 rows lost the distinction'),
       ('Hủy', 'Cancelled', E, "state 'cancelled' + stage 'Cancelled' (used: 15 rows)", '')]
for vi, en, s2, match, act in bst:
    addb(st_, vi, en, 'health.fieldservice.order.state + health.fieldservice.stage', s2, match, act)

sl = 'Dịch vụ tại (Service Location)'
addb(sl, 'Tại Nhà', 'At home', 'health.fieldservice.order.service_location', E, "'home' (hardcoded for ALL rows)")
addb(sl, 'Tại PK / Tại Phòng khám', 'At clinic', 'health.fieldservice.order.service_location', E, "'clinic' EXISTS",
     "Importer hardcoded 'home' - 21 clinic rows migrated WRONG; fix by re-mapping")
addb(sl, 'Telemedicine', 'Telemedicine', 'health.fieldservice.order.service_location', E, "'online' EXISTS", 'Not present in this export')

cr = 'Lý do hủy (Cancellation Reason)'
cr_names = [(_norm(i18n(x['name'])), i18n(x['name'])) for x in cancel_reasons]
def nearest_cr(en_txt):
    t2 = _norm(en_txt)
    for nrm, disp in cr_names:
        if t2 and (t2 in nrm or nrm in t2):
            return disp
    return ''
crs = [('Không có bác sĩ >2.5h', 'No doctor available > 2.5 hours', C, "nearest: 'Staff unavailable'"),
       ('Không có bác sĩ <2.5h', 'No doctor available < 2.5 hours', C, "nearest: 'Staff unavailable'"),
       ('Không có điều dưỡng >2.5h', 'No nurse available > 2.5 hours', C, "nearest: 'Staff unavailable'"),
       ('Không có điều dưỡng <2.5h', 'No nurse available < 2.5 hours', C, "nearest: 'Staff unavailable'"),
       ('Y tá được ưu tiên hiện không có sẵn', 'Preferred Nurse Not Available', C, "nearest: 'Staff unavailable'"),
       ('Nhập viện', 'Hospitalized', C, ''),
       ('Qua đời', 'Deceased', C, ''),
       ('Hết đơn thuốc', 'Prescription completed / expired', C, "nearest: 'Service no longer needed'"),
       ('Khác', 'Other', C, '')]
for vi, en, s2, match in crs:
    addb(cr, vi, en, 'health.booking.cancellation.reason (17 generic records exist; none match legacy list)', s2, match,
         'Create record; today the value sits only in free-text cancellation_notes')

ci = 'City'
addb(ci, 'Hà Nội', 'Hanoi', 'health.catchment.province', E, "'Hà Nội' (code 01)")
addb(ci, 'TPHCM', 'Ho Chi Minh City', 'health.catchment.province', E, "'TPHCM' (code 02)")

pay = 'Thanh toán (Payment Status)'
addb(pay, 'Chưa thanh toán', 'Unpaid', 'account.move.payment_state', E, "'not_paid'", 'Not migrated (invoices not created)')
addb(pay, 'Đã thanh toán', 'Paid', 'account.move.payment_state', E, "'paid'", "Payment txn created with status 'collected'")
addb(pay, 'Thanh toán một phần', 'Partial Payment', 'account.move.payment_state', E, "'partial'", 'Not present in export')

pm = 'Method of Payment'
addb(pm, 'Tiền mặt', 'Cash', 'health.payment.transaction.payment_method', E, "'cash' (used: 342 rows)")
addb(pm, 'Chuyển khoản ngân hàng', 'Bank Transfer', 'health.payment.transaction.payment_method', E, "'bank_transfer' (used: 550 rows)")
addb(pm, 'Trả từ TT trước', 'Paid from prepayment balance', 'health.payment.transaction.payment_method', E, "'prepaid' EXISTS",
     "Importer had no mapping - 66 rows fell back to 'other'; remap to 'prepaid'")
addb(pm, 'Thanh toán trước', 'Prepaid / Pay in advance', 'health.payment.transaction.payment_method', E, "'prepaid' EXISTS",
     "8 rows - same remap")

ip = 'In đơn (Invoice Print)'
addb(ip, 'In đơn lần 1', 'First invoice print', '(no target)', D, '', 'No print-count concept; e-invoice handled by health_redinvoice - decide if needed')
addb(ip, 'In đơn cuối cùng', 'Final invoice print', '(no target)', D, '', 'Same as above')

fa = 'Facility / Code'
fac_codes = {x['code'] for x in facilities}
for code, name, prov in [('01_01_00', 'HN_PK_PKGDVU', 'Hà Nội'), ('01_01_01', 'HN_NS_Minh_Khai', 'Hà Nội'),
                         ('01_01_02', 'HN_NS_Hà_Đông', 'Hà Nội'), ('02_01_00', 'TPHCM_PK_CSTNVU', 'TPHCM'),
                         ('02_01_01', 'TPHCM_NS_Thủ_Đức', 'TPHCM'), ('02_01_02', 'TPHCM_NS_Tân_Thuận', 'TPHCM')]:
    addb(fa, name, 'code %s (%s)' % (code, prov), 'health.facility', C, "only 2 facilities exist: 'Hanoi Facility' (HAN), 'HCM Facility' (HCM)",
         'Create facility with the official code; migrated bookings all point at the 2 city-level facilities')

def build_lookup_sheet(name, rows_):
    ws = out.create_sheet(name)
    ws.append(['Lookup List', 'Legacy Value (VI)', 'Legacy Value (EN)', 'Target Model / Field',
               'Status', 'Existing Match', 'Action / Note'])
    style_header(ws)
    ws.freeze_panes = 'A2'
    last = None
    for (ln, vi, en, tgt, s2, match, act) in rows_:
        ws.append([ln if ln != last else '', vi, en, tgt, {'EXISTS': 'EXISTS', 'CREATE': 'CREATE', 'DECISION': 'DECISION NEEDED'}[s2], match, act])
        last = ln
        r_ = ws.max_row
        cat = {'EXISTS': 'green', 'CREATE': 'orange', 'DECISION': 'red'}[s2]
        for cc in range(2, 8):
            ws.cell(row=r_, column=cc).alignment = WRAP
        paint(ws.cell(row=r_, column=5), cat)
        paint(ws.cell(row=r_, column=2), cat)
        ws.cell(row=r_, column=1).font = F_SECT
    for i, w in enumerate([30, 26, 26, 42, 16, 40, 55], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws

build_lookup_sheet('Lookup - Contact Table', ct_lookup)
build_lookup_sheet('Lookup - Booking Table', bk_lookup)

# ---------- summary sheet (built last, inserted first)
ws = out.create_sheet('Summary', 0)
ws.sheet_view.showGridLines = False
row = 1
def W(text, font=F_PLAIN, col=1, fill=None):
    global row
    c = ws.cell(row=row, column=col, value=text)
    c.font = font
    if fill:
        c.fill = fill
    row += 1
    return c

W('Pancake -> Health19 Migration Audit Report', F_TITLE)
W('Generated 2026-07-21 - verified against the LIVE vietuat database (VietUcUAT)', F_GREY.copy(italic=True) if hasattr(F_GREY, 'copy') else F_PLAIN)
row += 1
W('Colour legend', F_SECT)
for cat, label in [('green', 'Migrated as-is (value unchanged)'),
                   ('blue', 'Upsert key (verbatim identifier, verified in DB)'),
                   ('yellow', 'Migrated but TRANSFORMED to fit the system (mapped / parsed / merged)'),
                   ('red', 'LOST or FAILED - could not be migrated, or was dropped/overwritten'),
                   ('grey', 'Column not migrated (by design / no target field)'),
                   ('orange', 'Lookup value missing in the system - needs to be created')]:
    c = ws.cell(row=row, column=1, value='')
    c.fill, c.font = STYLE[cat]
    ws.cell(row=row, column=2, value=label).font = F_PLAIN
    row += 1
row += 1

W('Reconciliation vs vietuat (row-by-row, re-computed legacy keys)', F_SECT)
uniq_clients = len({str(rec_of(bk_hdr, r).get('Client ID') or '').strip() for r in bk_rows} - {''})
recon = [
    ('Booking rows in Booking_mig.xlsx', len(bk_rows), ''),
    ('Bookings found in vietuat (legacy_booking_ref)', len(bk_rows) - len(missing_refs),
     'MISSING: %d' % len(missing_refs) if missing_refs else 'ALL LANDED'),
    ('Unique clients in booking file', uniq_clients, ''),
    ('Clients found in vietuat (legacy_client_code)', len(db_client_codes), ''),
    ('Contact rows in Contact_mig.xlsx', len(ct_rows), ''),
    ('Leads found in vietuat (legacy_contact_guid)', len(ct_rows) - len(missing_guids),
     'MISSING: %d' % len(missing_guids) if missing_guids else 'ALL LANDED'),
    ('Sale orders created (origin BKG-*)', len(sale_orders), '%d booking rows had no sale order' % len(no_so) if no_so else ''),
    ('Payment transactions created', len(payments), '%d rows with money had no payment' % len(no_pay) if no_pay else ''),
    ('Auto-created service types / products', '%d / %d' % (len([x for x in xref if x['target_model'] == 'health.service.type']),
                                                           len([x for x in xref if x['target_model'] == 'product.product'])), 'ALL at price 0'),
    ('Auto-created staff stubs', len([x for x in xref if x['target_model'] == 'hr.employee' and x['auto_created'] == 't']), ''),
]
ws.cell(row=row, column=1, value='Metric').font = F_HDR; ws.cell(row=row, column=1).fill = FILL_HDR
ws.cell(row=row, column=2, value='Count').font = F_HDR; ws.cell(row=row, column=2).fill = FILL_HDR
ws.cell(row=row, column=3, value='Note').font = F_HDR; ws.cell(row=row, column=3).fill = FILL_HDR
row += 1
for m, v, note in recon:
    ws.cell(row=row, column=1, value=m).font = F_PLAIN
    ws.cell(row=row, column=2, value=v).font = F_PLAIN
    nc = ws.cell(row=row, column=3, value=note)
    nc.font = F_RED if note.startswith(('MISSING', 'ALL at')) or 'no ' in note else F_GREEN
    row += 1
row += 1

W('Top findings - data LOST or materially changed', F_SECT)
findings = [
    '1. PHONES: the export truncated every phone to 3-4 digits. None validated; the system now holds RANDOM PLACEHOLDER numbers on all migrated clients and leads. Real numbers must be re-imported from an unredacted export.',
    '2. MONEY: SubTotal / Discount / GrandTotal / surcharge (phu_phi) / prepayment were NOT migrated. All 166 products and 68 service types were created at price 0 and every sale line is 0 VND. Only the Cash+Transfer sum went in as one payment transaction (843 payments).',
    "3. SERVICE LOCATION: importer hardcoded 'home'. 21 'Tại Phòng khám' (at-clinic) bookings are wrong in the system ('clinic' value exists).",
    "4. CLIENT SOURCE: SourceofClient (booking) was dropped - booking_source is hardcoded 'phone'. On contacts only Hotline/Zalo/'Khách hàng cũ' were mapped: 107 leads have an EMPTY lead source.",
    "5. GENDER: contact-file gender was never written (crm.lead has no gender field; 289 values dropped). 'Không xác định' on bookings (15) dropped to blank.",
    "6. STATUS GRANULARITY: 'Dời lịch' (rescheduled, 5 rows) collapsed into Confirmed. Payment method 'Trả từ TT trước' (66) and 'Thanh toán trước' (8) fell back to 'other' although 'prepaid' exists.",
    '7. CANCEL/REJECT REASONS: kept as free text only; not linked to the structured lookup models (see Lookup sheets - most legacy reasons need creating).',
    '8. NOT MIGRATED BY DESIGN (grey): ad/channel tracking (page_id, ad_id, conversation_id, Facebook Link, psid), addresses beyond one street line, ethnicity, parking fee, performer, distances (re-geocoded), timestamps/owners (CreatedBy/On, ModifiedBy/On), invoice-print flags.',
]
for f in findings:
    c = ws.cell(row=row, column=1, value=f)
    c.font = F_PLAIN; c.alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.row_dimensions[row].height = 30
    row += 1
row += 1

W('Lookup coverage (Lookup.xlsx vs system)', F_SECT)
def lkstats(rows_):
    e = sum(1 for x in rows_ if x[4] == 'EXISTS'); c = sum(1 for x in rows_ if x[4] == 'CREATE')
    d = sum(1 for x in rows_ if x[4] == 'DECISION')
    return e, c, d
e1, c1_, d1 = lkstats(ct_lookup); e2, c2_, d2 = lkstats(bk_lookup)
W('Contact Table: %d exist, %d need creation, %d need a decision' % (e1, c1_, d1))
W('Booking Table: %d exist, %d need creation, %d need a decision' % (e2, c2_, d2))
W('NOTE: the Lookup codes file itself was NEVER migrated as data. The importer used small hard-coded maps; everything else in Lookup.xlsx exists only where marked green.', F_SECT)
row += 1
W('Sheets: Mapping - * = column-by-column field mapping | * Data = full colour-coded mirror of the source file | * Issues = every lost/changed cell with explanation | Lookup - * = per-value existence check', F_PLAIN)
ws.column_dimensions['A'].width = 60
ws.column_dimensions['B'].width = 46
ws.column_dimensions['C'].width = 40

# order sheets
order = ['Summary', 'Mapping - Booking', 'Mapping - Contact', 'Booking Data', 'Contact Data',
         'Booking Issues', 'Contact Issues', 'Lookup - Contact Table', 'Lookup - Booking Table']
out._sheets = [out[n] for n in order]

out.save(OUT)
print('saved', OUT)
print('booking cells:', bk_counts)
print('contact cells:', ct_counts)
print('booking issues:', len(bk_issues), 'contact issues:', len(ct_issues))
print('missing refs:', len(missing_refs), 'missing guids:', len(missing_guids), 'no SO:', len(no_so), 'no pay:', len(no_pay))
