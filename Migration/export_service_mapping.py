#!/usr/bin/env python3
"""Export the reviewed service mapping to JSON for the migration to apply.

Reads Service_Mapping_For_Review.xlsx. Where ops filled in the "CONFIRM"
columns those win; otherwise the proposed mapping stands.

Output: Migration/service_mapping.json
  groups   : the price-list service_name groups (the master service catalogue),
             each with the service_type that decides the platform category
  segments : every legacy dich_vu string -> kind / item_code / group
"""
import json
import os
import unicodedata
import re

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'Service_Mapping_For_Review.xlsx')
PRICELIST = os.path.join(HERE, 'Price_List_Adjustments_Rules_20250917 V3.2 Scrubbed.xlsx')
DST = os.path.join(HERE, 'service_mapping.json')

# price-list service_type -> health.service.type.category key
STYPE_TO_CATEGORY = {'DV Bs': 'consultation', 'DV ĐD': 'nursing_care'}

# price-list service_name group -> crm.lead.service_interest key
GROUP_TO_INTEREST = {
    'Khám & Điều trị': 'consultation',
    'Chăm sóc VT': 'wound_care',
    'Truyền & Tiêm': 'iv_infusion',
    'Sonde': 'catheter',
    'Hỗ trợ hô hấp': 'sputum_care',
    'Thụt tháo': 'enema',
    'CSGN': 'palliative',
    'CSCN': 'personal_care',
    'Mẫu bệnh phẩm': 'lab_test',
    'Dịch vụ yêu cầu khác': 'other',
    'Ảnh y tế & Thủ thuật': 'imaging_procedures',
    'Điều dưỡng Khác': 'nursing_other',
}


def norm(s):
    """Must stay identical to migration_runner._norm_key, or the importer will
    not recognise the strings this file exports."""
    s = unicodedata.normalize('NFC', str(s or '')).strip().lower()
    s = re.sub(r'\s+', ' ', s)
    s = re.sub(r'\s*([<>])\s*', r'\1', s)
    return s.strip(' .,;')


def load_groups():
    wb = openpyxl.load_workbook(PRICELIST, read_only=True, data_only=True)
    groups = {}
    for sh in ('Hanoi_Base_Clean_v3', 'HCMC_Base_Clean_v3'):
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h or '') for h in rows[0]]
        i = {h: n for n, h in enumerate(hdr)}
        for r in rows[1:]:
            if not any(r):
                continue
            vi = str(r[i['service_name_vi']] or '').strip()
            if not vi:
                continue
            g = groups.setdefault(vi, {
                'vi': vi,
                'en': str(r[i['service_name_en']] or '').strip(),
                'service_types': set(),
                'items': 0,
            })
            if not g['en'] and r[i['service_name_en']]:
                g['en'] = str(r[i['service_name_en']]).strip()
            g['service_types'].add(str(r[i['service_type']] or '').strip())
            g['items'] += 1
    out = []
    for vi, g in groups.items():
        # a group belongs to the doctor catalogue only if every item does
        stypes = {s for s in g['service_types'] if s}
        stype = 'DV Bs' if stypes == {'DV Bs'} else 'DV ĐD'
        out.append({
            'vi': vi,
            'en': g['en'] or vi,
            'service_type': stype,
            'category': STYPE_TO_CATEGORY[stype],
            'interest': GROUP_TO_INTEREST.get(vi, ''),
            'items': g['items'],
        })
    return out


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb['Service Mapping']
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    segments = []
    for r in rows:
        seg = r[0]
        if not seg:
            continue
        confirmed_code = str(r[12] or '').strip()
        confirmed_group = str(r[13] or '').strip()
        segments.append({
            'text': seg,
            'norm': norm(seg),
            'bookings': r[1] or 0,
            'kind': r[3],                      # SERVICE / ADJUSTMENT / LAB TEST
            'item_code': confirmed_code or str(r[4] or '').strip(),
            'service_type': str(r[5] or '').strip(),
            'group_vi': confirmed_group or str(r[6] or '').strip(),
            'confidence': r[10],
        })
    groups = load_groups()
    payload = {'groups': groups, 'segments': segments}
    with open(DST, 'w', encoding='utf-8') as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=1)
    kinds = {}
    for s in segments:
        kinds[s['kind']] = kinds.get(s['kind'], 0) + 1
    coded = sum(1 for s in segments if s['item_code'])
    print('Wrote %s' % DST)
    print('  groups   : %d' % len(groups))
    print('  segments : %d  %s' % (len(segments), kinds))
    print('  with a price-list item_code: %d' % coded)


if __name__ == '__main__':
    main()
