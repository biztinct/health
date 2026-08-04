#!/usr/bin/env python3
"""Service mapping sheet — every legacy booking service string vs the price list.

The legacy export records services as free text in `dich_vu` (semicolon
separated). The price list is the master: it defines the item, its service_type
(DV Bs / DV ĐD = the legacy Contact_dich_vu_danh_muc) and its service_name group
(= the legacy Contact_dich_vu_quan_tam). This script proposes, for every distinct
legacy string, which price-list item it is — so ops can confirm or correct it
before the migration derives categories from it.

Each row is classified first:
  SERVICE     — a real billable service, should resolve to a price-list item
  ADJUSTMENT  — a surcharge (travel, after-hours, weekend). These are priced by
                the Adjustments rules, not sold as catalogue items. They stay as
                invoice lines so revenue reconciles, but must never become a
                service type or appointment type.

Output: Migration/Service_Mapping_For_Review.xlsx
"""
import collections
import difflib
import os
import re
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
PRICELIST = os.path.join(HERE, 'Price_List_Adjustments_Rules_20250917 V3.2 Scrubbed.xlsx')
BOOKINGS = os.path.join(HERE, 'Booking_mig.xlsx')
DST = os.path.join(HERE, 'Service_Mapping_For_Review.xlsx')

GREEN = PatternFill('solid', fgColor='C6EFCE')
AMBER = PatternFill('solid', fgColor='FFE0B2')
RED = PatternFill('solid', fgColor='FFC7CE')
GREY = PatternFill('solid', fgColor='E7E6E6')
HEAD = PatternFill('solid', fgColor='1F3864')
WRAP = Alignment(wrap_text=True, vertical='top')

# A surcharge is recognised from the Adjustments sheets plus the wording the
# operators actually used in the booking file.
ADJ_WORDS = re.compile(
    r'phụ phí|ngoài giờ|cuối tuần|khoảng cách|đi lại|phát sinh', re.I)

# Individual laboratory analytes. The price list has no line item for any of
# them — it only carries a specimen-collection fee — so they are called out
# separately rather than being forced onto a service item.
LAB_WORDS = re.compile(
    r'\b(crp|albumin|albumine|asat|alat|sgot|sgpt|ast|alt|ck|ckmb|ck ?- ?mb|'
    r'procalcitonin|pct|bnp|probnp|nt-probnp|ggt|gamma ?gt|calcium|calci|tsh|'
    r'ft3|ft4|free ?t4|troponin|troponine|dengue|ns1|igg|igm|gram|lactate|'
    r'egfr|dimer|dimères|bilirubin|cortisol|ferritin|transferin|magnesium|'
    r'aptt|tck|creatinine|ure|glucose|hba1c|cholesterol|triglyceride|'
    r'acid ?uric|acid ?lactic|protein|prothrombin|toxocara|strongyloides|'
    r'gnathostoma|kháng sinh đồ|cấy|nhuộm soi|điện giải|huyết thanh|'
    r'sắt huyết thanh|pt \(tp)', re.I)

# Keyword -> price-list service_name group. Derived from the group names
# themselves, so a string that cannot be pinned to one item can still be placed
# in the right group (which is all the category derivation needs).
GROUP_RULES = [
    (r'theo yêu cầu|yêu cầu khác|dịch vụ khác', 'Dịch vụ yêu cầu khác'),
    (r'thay băng|cắt chỉ|vết thương|vt\b|nhiễm trùng|loét', 'Chăm sóc VT'),
    (r'tiêm|truyền|chai|tĩnh mạch|thuốc', 'Truyền & Tiêm'),
    (r'sonde|thông tiểu|bàng quang|dạ dày|tiểu lưu', 'Sonde'),
    (r'thụt tháo', 'Thụt tháo'),
    (r'hút đờm|khí dung|hô hấp|vỗ rung', 'Hỗ trợ hô hấp'),
    (r'giảm nhẹ|csgn', 'CSGN'),
    (r'cá nhân|cscn|tắm|vệ sinh', 'CSCN'),
    (r'xét nghiệm|mẫu|máu|creatinine|ure|ast|alt|glucose|điện giải|'
     r'huyết học|sinh hóa|nước tiểu|đường huyết|hba1c|cholesterol|'
     r'triglyceride|acid uric|gót chân|tổng phân tích', 'Mẫu bệnh phẩm'),
    (r'siêu âm|điện tim|điện tâm đồ|chọc dịch|x-quang', 'Ảnh y tế & Thủ thuật'),
    (r'khám|bác sĩ|bs\b|tư vấn|nội khoa', 'Khám & Điều trị'),
    (r'điều dưỡng', 'Điều dưỡng Khác'),
]


def norm(s):
    s = unicodedata.normalize('NFC', str(s or '')).strip().lower()
    s = re.sub(r'\s+', ' ', s)
    return s.strip(' .,;:')


def deaccent(s):
    s = unicodedata.normalize('NFD', norm(s))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return s.replace('đ', 'd')


def load_pricelist():
    wb = openpyxl.load_workbook(PRICELIST, read_only=True, data_only=True)
    items, groups, adj_names = [], {}, set()
    for sh in ('Hanoi_Base_Clean_v3', 'HCMC_Base_Clean_v3'):
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h or '') for h in rows[0]]
        i = {h: n for n, h in enumerate(hdr)}
        for r in rows[1:]:
            if not any(r):
                continue
            desc = str(r[i['item_description']] or '').strip()
            if not desc:
                continue
            vi = str(r[i['service_name_vi']] or '').strip()
            en = str(r[i['service_name_en']] or '').strip()
            items.append({
                'code': str(r[i['item_code']] or '').strip(),
                'province': str(r[i['province']] or '').strip(),
                'desc': desc,
                'stype': str(r[i['service_type']] or '').strip(),
                'scat': str(r[i['service_category']] or '').strip(),
                'group_vi': vi,
                'group_en': en,
                'fee': r[i['base_fee_vnd']],
            })
            if vi:
                groups.setdefault(vi, en)
    for sh in ('Hanoi_Adjustments_Clean', 'HCMC_Adjustments_Clean'):
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h or '') for h in rows[0]]
        i = {h: n for n, h in enumerate(hdr)}
        col = 'Adjustment Applies to this item_name'
        for r in rows[1:]:
            if any(r) and col in i:
                adj_names.add(norm(r[i[col]]))
    return items, groups, adj_names


def load_segments():
    wb = openpyxl.load_workbook(BOOKINGS, read_only=True, data_only=True)
    ws = wb.active
    hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
    i = {h: n for n, h in enumerate(hdr) if h}
    counts, as_first = collections.Counter(), collections.Counter()
    for r in ws.iter_rows(min_row=2, values_only=True):
        segs = [x.strip() for x in str(r[i['dich_vu']] or '').split(';') if x.strip()]
        for n, s in enumerate(segs):
            counts[s] += 1
            if n == 0:
                as_first[s] += 1
    return counts, as_first


def best_item(seg, items):
    """Closest price-list item by description, accent-insensitive."""
    target = deaccent(seg)
    best, score = None, 0.0
    for it in items:
        s = difflib.SequenceMatcher(None, target, deaccent(it['desc'])).ratio()
        if s > score:
            best, score = it, s
    return best, score


def guess_group(seg, groups):
    d = deaccent(seg)
    for pattern, group_vi in GROUP_RULES:
        if re.search(deaccent(pattern).replace(r'\b', ''), d):
            return group_vi, groups.get(group_vi, '')
    return '', ''


def build():
    items, groups, adj_names = load_pricelist()
    counts, as_first = load_segments()

    rows = []
    for seg, n in counts.most_common():
        is_adj = bool(ADJ_WORDS.search(seg)) or norm(seg) in adj_names
        it, score = best_item(seg, items)
        gvi, gen = guess_group(seg, groups)
        is_lab = (not is_adj and score < 0.90
                  and bool(LAB_WORDS.search(deaccent(seg)) or LAB_WORDS.search(seg)))
        if is_lab:
            kind = 'LAB TEST'
            code = ''
            stype = 'DV ĐD'
            gvi, gen = 'Mẫu bệnh phẩm', groups.get('Mẫu bệnh phẩm', 'Specimen Sample')
            conf = 'LAB — not in price list'
            note = ('An individual laboratory analyte. The price list has no line item '
                    'for it (only a specimen-collection fee), so it cannot resolve to a '
                    'code. Decide: add these to the price list, or keep them as '
                    'pass-through lab charges under Specimen Sample.')
        elif is_adj:
            kind = 'ADJUSTMENT'
            code = ''
            stype = ''
            gvi = gen = ''
            conf = 'N/A — surcharge'
            note = ('Priced by the Adjustments rules, not a catalogue item. Keep as an '
                    'invoice line so revenue reconciles; must not become a service '
                    'type or appointment type.')
        else:
            kind = 'SERVICE'
            if score >= 0.90:
                conf, code, stype = 'HIGH', it['code'], it['stype']
                gvi, gen = it['group_vi'], it['group_en']
                note = 'Near-exact match on the price-list item description.'
            elif score >= 0.62 and it and gvi and it['group_vi'] == gvi:
                conf, code, stype = 'MEDIUM', it['code'], it['stype']
                gvi, gen = it['group_vi'], it['group_en']
                note = ('Wording differs but the item and the keyword rules agree on '
                        'the same group.')
            elif gvi:
                conf, code = 'GROUP ONLY', ''
                stype = 'DV Bs' if gvi in ('Khám & Điều trị',) else 'DV ĐD'
                note = ('No single item is a confident match, but the wording places it '
                        'in this group — enough to derive the category. Pick the exact '
                        'item if you want per-item pricing.')
            else:
                conf, code, stype = 'NO MATCH', '', ''
                note = 'Could not be placed. Needs an ops decision.'
        rows.append({
            'seg': seg, 'n': n, 'first': as_first.get(seg, 0), 'kind': kind,
            'code': code, 'stype': stype, 'gvi': gvi, 'gen': gen,
            'cand': it['desc'] if it else '', 'cand_code': it['code'] if it else '',
            'score': round(score, 2), 'conf': conf, 'note': note,
        })

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Service Mapping'
    headers = ['Legacy Service Text (dich_vu)', 'Bookings', 'As Primary Service',
               'Type', 'Proposed item_code', 'service_type (danh mục)',
               'service_name group (quan tâm)', 'Group (EN)',
               'Closest Price-List Item', 'Match', 'Confidence',
               'Why / What to do',
               '✔ CONFIRM item_code', '✔ CONFIRM group', 'Ops Notes']
    widths = [52, 10, 12, 13, 16, 18, 24, 24, 46, 8, 14, 54, 18, 22, 26]
    bold = Font(bold=True, color='FFFFFF')
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = bold
        c.fill = HEAD
        c.alignment = WRAP
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = 'A2'

    r = 2
    for d in rows:
        vals = [d['seg'], d['n'], d['first'], d['kind'], d['code'], d['stype'],
                d['gvi'], d['gen'], d['cand'], d['score'], d['conf'], d['note'],
                '', '', '']
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v).alignment = WRAP
        fill = {'HIGH': GREEN, 'MEDIUM': GREEN, 'GROUP ONLY': AMBER,
                'LAB — not in price list': AMBER, 'NO MATCH': RED}.get(d['conf'], GREY)
        ws.cell(row=r, column=11).fill = fill
        ws.cell(row=r, column=4).fill = (
            GREY if d['kind'] == 'ADJUSTMENT'
            else AMBER if d['kind'] == 'LAB TEST' else GREEN)
        r += 1

    # ---------------------------------------------------------------- summary
    ws2 = wb.create_sheet('Summary', 0)
    ws2.column_dimensions['A'].width = 62
    ws2.column_dimensions['B'].width = 14
    ws2.column_dimensions['C'].width = 14
    tally = collections.Counter(d['conf'] for d in rows)
    kinds = collections.Counter(d['kind'] for d in rows)
    occ = collections.Counter()
    for d in rows:
        occ[d['conf']] += d['n']
    lines = [
        ('Service mapping for review — legacy booking text vs the price list', None, None),
        ('', None, None),
        ('The price list is the master. Confirm the proposed item / group for each legacy '
         'service string; the migration then derives the service category '
         '(Contact_dich_vu_danh_muc) and the contact service interest '
         '(Contact_dich_vu_quan_tam) from it instead of guessing.', None, None),
        ('', None, None),
        ('', 'Distinct strings', 'Booking lines'),
        ('Real services', kinds.get('SERVICE', 0),
         sum(d['n'] for d in rows if d['kind'] == 'SERVICE')),
        ('Surcharges (priced by Adjustment rules, not catalogue items)',
         kinds.get('ADJUSTMENT', 0),
         sum(d['n'] for d in rows if d['kind'] == 'ADJUSTMENT')),
        ('Laboratory analytes (no line item exists in the price list)',
         kinds.get('LAB TEST', 0),
         sum(d['n'] for d in rows if d['kind'] == 'LAB TEST')),
        ('', None, None),
        ('Confidence of the proposed mapping', 'Distinct strings', 'Booking lines'),
        ('HIGH — near-exact item match', tally.get('HIGH', 0), occ.get('HIGH', 0)),
        ('MEDIUM — item and keywords agree on the group', tally.get('MEDIUM', 0),
         occ.get('MEDIUM', 0)),
        ('GROUP ONLY — group is clear, exact item is not', tally.get('GROUP ONLY', 0),
         occ.get('GROUP ONLY', 0)),
        ('LAB — analyte, no price-list item', tally.get('LAB — not in price list', 0),
         occ.get('LAB — not in price list', 0)),
        ('NO MATCH — needs an ops decision', tally.get('NO MATCH', 0),
         occ.get('NO MATCH', 0)),
        ('', None, None),
        ('How to use this sheet: work down the Service Mapping tab, hardest first '
         '(NO MATCH, then GROUP ONLY). Put the correct code in "CONFIRM item_code" or '
         'the correct group in "CONFIRM group". Anything left blank keeps the proposal.',
         None, None),
    ]
    rr = 1
    for a, b, c in lines:
        ca = ws2.cell(row=rr, column=1, value=a)
        ca.alignment = WRAP
        if b is not None:
            ws2.cell(row=rr, column=2, value=b).font = Font(bold=True)
        if c is not None:
            ws2.cell(row=rr, column=3, value=c).font = Font(bold=True)
        if a and b is None and c is None and rr in (1,):
            ca.font = Font(bold=True, size=12)
        rr += 1
    wb.save(DST)
    print('Wrote', DST)
    print('services=%d adjustments=%d | %s' % (
        kinds.get('SERVICE', 0), kinds.get('ADJUSTMENT', 0), dict(tally)))


if __name__ == '__main__':
    build()
